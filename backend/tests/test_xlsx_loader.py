import importlib.util
import tempfile
import unittest
from pathlib import Path

from app.agent.ingestion_agent import plan_from_probe, probe_document
from app.rag.file_types import detect_file_type
from app.rag.loaders.factory import get_loader
from app.rag.loaders.xlsx_loader import XlsxLoader
from app.services.rag_service import (
    _format_citation_reference,
    _hit_to_citation,
)


@unittest.skipUnless(
    importlib.util.find_spec("openpyxl") is not None,
    "openpyxl is not installed",
)
class XlsxLoaderTests(unittest.TestCase):
    def test_extracts_merged_title_headers_rows_and_formulas(self) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "quality-limits.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Limits"
            sheet.merge_cells("A1:D1")
            sheet["A1"] = "Meat quality limits"
            sheet["A1"].font = Font(bold=True)
            sheet.append(["Product", "Metric", "Base", "Calculated"])
            for cell in sheet[2]:
                cell.font = Font(bold=True)
            sheet.append(["Fresh meat", "TVBN", 15, "=C3*2"])
            sheet.append(["Chilled meat", "TVBN", 15, "=C4*2"])
            hidden = workbook.create_sheet("HiddenData")
            hidden.sheet_state = "hidden"
            hidden.append(["Secret", "Value"])
            hidden.append(["Internal", 99])
            workbook.save(path)

            blocks = XlsxLoader().load(path)

        context_blocks = [
            block for block in blocks if block.block_type == "sheet_context"
        ]
        markdown_blocks = [
            block for block in blocks if block.block_type == "table_markdown"
        ]
        row_blocks = [
            block for block in blocks if block.block_type == "table_row_fact"
        ]

        self.assertEqual(len(context_blocks), 1)
        self.assertEqual(context_blocks[0].text, "Meat quality limits")
        self.assertEqual(len(markdown_blocks), 1)
        self.assertIn("Product", markdown_blocks[0].text)
        self.assertEqual(len(row_blocks), 2)
        self.assertIn("TVBN", row_blocks[0].text)
        self.assertIn("=C3*2", row_blocks[0].text)
        self.assertEqual(row_blocks[0].metadata["sheet"], "Limits")
        self.assertEqual(row_blocks[0].metadata["cell_range"], "A3:D3")
        self.assertEqual(row_blocks[0].metadata["formula_count"], 1)
        self.assertNotIn("HiddenData", {block.metadata.get("sheet") for block in blocks})

    def test_large_table_markdown_is_segmented(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "large.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Rows"
            sheet.append(["Item", "Value"])
            for index in range(45):
                sheet.append([f"Item {index}", index])
            workbook.save(path)

            blocks = XlsxLoader().load(path)

        markdown_blocks = [
            block for block in blocks if block.block_type == "table_markdown"
        ]
        row_blocks = [
            block for block in blocks if block.block_type == "table_row_fact"
        ]
        self.assertEqual(len(markdown_blocks), 2)
        self.assertEqual(len(row_blocks), 45)
        self.assertEqual(markdown_blocks[0].metadata["table_segment_count"], 2)

    def test_factory_probe_plan_and_citation_support_xlsx(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "demo.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Standards"
            sheet.append(["Metric", "Limit"])
            sheet.append(["TVBN", "15 mg/100g"])
            workbook.save(path)

            probe = probe_document(path, "xlsx")
            plan = plan_from_probe(probe)

        self.assertEqual(detect_file_type("demo.xlsx"), "xlsx")
        self.assertIsInstance(get_loader("xlsx"), XlsxLoader)
        self.assertEqual(probe.page_count, 1)
        self.assertGreater(probe.text_chars, 0)
        self.assertEqual(plan.primary_reader, "xlsx_loader")
        self.assertFalse(plan.use_ocr)
        self.assertFalse(plan.use_llm_preprocessing)

        citation = _hit_to_citation(
            {
                "id": "doc:3:chunk:4",
                "score": 0.91,
                "text": "spreadsheet evidence",
                "metadata": {
                    "document_id": 3,
                    "filename": "limits.xlsx",
                    "sheet": "Standards",
                    "cell_range": "A12:D12",
                    "chunk_type": "table_row_fact",
                },
            }
        )
        self.assertEqual(citation.sheet, "Standards")
        self.assertEqual(
            _format_citation_reference(citation),
            "limits.xlsx，工作表 Standards / 单元格 A12:D12",
        )


if __name__ == "__main__":
    unittest.main()
