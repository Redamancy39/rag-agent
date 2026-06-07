import importlib.util
import tempfile
import unittest
from pathlib import Path

from app.agent.ingestion_agent import plan_from_probe, probe_document
from app.rag.file_types import detect_file_type
from app.rag.loaders.factory import get_loader
from app.rag.loaders.pptx_loader import PptxLoader
from app.services.rag_service import (
    _format_citation_reference,
    _hit_to_citation,
)


@unittest.skipUnless(
    importlib.util.find_spec("pptx") is not None,
    "python-pptx is not installed",
)
class PptxLoaderTests(unittest.TestCase):
    def test_extracts_slide_text_table_and_notes(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "quality-report.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "产品质量要求"
            slide.placeholders[1].text = "适用范围：鲜肉产品\n检测项目：挥发性盐基氮"
            table_shape = slide.shapes.add_table(
                rows=2,
                cols=3,
                left=Inches(1),
                top=Inches(3),
                width=Inches(8),
                height=Inches(1.5),
            )
            table = table_shape.table
            table.cell(0, 0).text = "产品"
            table.cell(0, 1).text = "项目"
            table.cell(0, 2).text = "限量"
            table.cell(1, 0).text = "鲜肉"
            table.cell(1, 1).text = "挥发性盐基氮"
            table.cell(1, 2).text = "15 mg/100g"
            slide.notes_slide.notes_text_frame.text = "讲解时说明检测方法和适用条件。"
            presentation.save(path)

            blocks = PptxLoader().load(path)

        slide_blocks = [
            block for block in blocks if block.block_type == "slide_text"
        ]
        row_blocks = [
            block for block in blocks if block.block_type == "table_row_fact"
        ]
        note_blocks = [
            block for block in blocks if block.block_type == "speaker_notes"
        ]

        self.assertEqual(len(slide_blocks), 1)
        self.assertEqual(slide_blocks[0].text.count("产品质量要求"), 1)
        self.assertIn("适用范围", slide_blocks[0].text)
        self.assertEqual(slide_blocks[0].metadata["slide"], 1)
        self.assertEqual(slide_blocks[0].metadata["document_format"], "pptx")
        self.assertEqual(len(row_blocks), 1)
        self.assertIn("挥发性盐基氮", row_blocks[0].text)
        self.assertIn("15 mg/100g", row_blocks[0].text)
        self.assertEqual(row_blocks[0].metadata["table_index"], 1)
        self.assertEqual(len(note_blocks), 1)
        self.assertIn("检测方法", note_blocks[0].text)

    def test_factory_detection_probe_and_plan_support_pptx(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "demo.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Overview"
            slide.placeholders[1].text = "A readable presentation body."
            presentation.save(path)

            probe = probe_document(path, "pptx")
            plan = plan_from_probe(probe)

        self.assertEqual(detect_file_type("demo.pptx"), "pptx")
        self.assertIsInstance(get_loader("pptx"), PptxLoader)
        self.assertEqual(probe.page_count, 1)
        self.assertGreater(probe.text_chars, 0)
        self.assertEqual(plan.primary_reader, "pptx_loader")
        self.assertFalse(plan.use_ocr)

    def test_pptx_citation_uses_slide_location(self) -> None:
        citation = _hit_to_citation(
            {
                "id": "doc:2:chunk:1",
                "score": 0.9,
                "text": "slide evidence",
                "metadata": {
                    "document_id": 2,
                    "filename": "report.pptx",
                    "slide": 4,
                    "section_path": "风险控制",
                    "chunk_type": "slide_text",
                },
            }
        )

        self.assertEqual(citation.slide, 4)
        self.assertEqual(
            _format_citation_reference(citation),
            "report.pptx，第 4 张幻灯片 / 章节 风险控制",
        )


if __name__ == "__main__":
    unittest.main()
