import importlib.util
import tempfile
import unittest
from pathlib import Path

from app.agent.ingestion_agent import plan_from_probe, probe_document
from app.rag.file_types import detect_file_type
from app.rag.loaders.docx_loader import DocxLoader
from app.rag.loaders.factory import get_loader
from app.services.rag_service import (
    _format_citation_reference,
    _hit_to_citation,
)


@unittest.skipUnless(
    importlib.util.find_spec("docx") is not None,
    "python-docx is not installed",
)
class DocxLoaderTests(unittest.TestCase):
    def test_extracts_sections_paragraphs_and_table_rows(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "quality-standard.docx"
            document = Document()
            document.add_heading("1 范围", level=1)
            document.add_paragraph("本文件规定了产品的质量要求。")
            document.add_paragraph("适用于企业内部质量检查。")
            document.add_heading("2 限量要求", level=1)
            document.add_paragraph("表1 产品限量")
            table = document.add_table(rows=1, cols=3)
            table.rows[0].cells[0].text = "产品"
            table.rows[0].cells[1].text = "项目"
            table.rows[0].cells[2].text = "限量"
            row = table.add_row().cells
            row[0].text = "鲜肉"
            row[1].text = "挥发性盐基氮"
            row[2].text = "15 mg/100g"
            document.save(path)

            blocks = DocxLoader().load(path)

        section_blocks = [
            block for block in blocks if block.block_type == "section_text"
        ]
        row_blocks = [
            block for block in blocks if block.block_type == "table_row_fact"
        ]
        markdown_blocks = [
            block for block in blocks if block.block_type == "table_markdown"
        ]

        self.assertEqual(len(section_blocks), 2)
        self.assertIn("本文件规定", section_blocks[0].text)
        self.assertIn("适用于", section_blocks[0].text)
        self.assertEqual(section_blocks[0].metadata["section_path"], "1 范围")
        self.assertEqual(section_blocks[1].metadata["section_path"], "2 限量要求")
        self.assertEqual(len(markdown_blocks), 1)
        self.assertEqual(len(row_blocks), 1)
        self.assertIn("挥发性盐基氮", row_blocks[0].text)
        self.assertIn("15 mg/100g", row_blocks[0].text)
        self.assertEqual(row_blocks[0].metadata["table_index"], 1)
        self.assertEqual(row_blocks[0].metadata["block_index"], 1)
        self.assertEqual(row_blocks[0].metadata["document_format"], "docx")

    def test_factory_detection_and_probe_support_docx(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "demo.docx"
            document = Document()
            document.add_heading("Overview", level=1)
            document.add_paragraph(
                "A readable DOCX paragraph for ingestion probing."
            )
            document.save(path)

            probe = probe_document(path, "docx")
            plan = plan_from_probe(probe)

        self.assertEqual(detect_file_type("demo.docx"), "docx")
        self.assertIsInstance(get_loader("docx"), DocxLoader)
        self.assertGreater(probe.text_chars, 0)
        self.assertEqual(plan.primary_reader, "docx_loader")
        self.assertFalse(plan.use_ocr)

    def test_docx_citation_uses_section_and_paragraph_location(self) -> None:
        citation = _hit_to_citation(
            {
                "id": "doc:1:chunk:2",
                "score": 0.88,
                "text": "source text",
                "metadata": {
                    "document_id": 1,
                    "filename": "policy.docx",
                    "section_path": "2 质量要求 / 2.1 限量",
                    "paragraph_index": 8,
                    "chunk_type": "section_text",
                },
            }
        )

        self.assertEqual(citation.section_path, "2 质量要求 / 2.1 限量")
        self.assertEqual(citation.paragraph_index, 8)
        self.assertEqual(
            _format_citation_reference(citation),
            "policy.docx，章节 2 质量要求 / 2.1 限量 / 段落 8",
        )


if __name__ == "__main__":
    unittest.main()
