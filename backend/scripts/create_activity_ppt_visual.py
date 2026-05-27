from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
ASSET_DIR = DOCS / "ppt_assets"
OUT = DOCS / "RAG_Agent_KB_System_activity_report_visual.pptx"

SOURCE_IMAGES = {
    "ui_upload": Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/e9802844-e358-478d-ad38-d397e9e3665d.png"),
    "ui_chat": Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/bbaccd31-c0ba-456b-b7a3-651a61206364.png"),
    "ui_progress": Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/68be821e-adaf-4b39-9e4f-b7c0b2f13426.png"),
    "case_scan": Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/fd132a56-daf2-42a7-80b7-bd9b1952ca74.png"),
    "case_watermark": Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/80fb6e19-cc2c-4c0b-8475-578db475a3f8.png"),
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(15, 23, 42)
INK = RGBColor(30, 41, 59)
GRAY = RGBColor(100, 116, 139)
LIGHT = RGBColor(247, 250, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(203, 213, 225)
TEAL = RGBColor(13, 124, 116)
MINT = RGBColor(223, 245, 241)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(229, 243, 255)
PINK = RGBColor(225, 29, 72)
ORANGE = RGBColor(245, 158, 11)
GREEN = RGBColor(22, 163, 74)
PURPLE = RGBColor(124, 58, 237)
RED = RGBColor(220, 38, 38)

FONT = "Microsoft YaHei"


def prepare_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    outputs: dict[str, Path] = {}
    for name, src in SOURCE_IMAGES.items():
        out = ASSET_DIR / f"{name}.png"
        if src.exists():
            with Image.open(src) as img:
                img = img.convert("RGB")
                # UI screenshots contain Windows activation text in the lower-right corner.
                # Crop a little from the bottom for cleaner slides.
                if name == "ui_upload":
                    img = img.crop((0, 0, img.width, int(img.height * 0.90)))
                elif name == "ui_chat":
                    img = img.crop((0, 0, img.width, int(img.height * 0.92)))
                elif name == "ui_progress":
                    img = img.crop((0, 0, img.width, int(img.height * 0.98)))
                img.save(out, quality=95)
        outputs[name] = out
    return outputs


def set_bg(slide, color=LIGHT):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    bold: bool = False,
    color=NAVY,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_lines(slide, lines, x, y, w, h, *, size=14, color=INK, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {line}" if bullet else line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
    return box


def add_title(slide, title: str, subtitle: str | None = None, tag: str | None = None):
    if tag:
        pill(slide, tag, 0.58, 0.36, 1.35, 0.34, fill=CYAN, color=BLUE, size=10)
    add_text(slide, title, 0.58, 0.78, 9.2, 0.48, size=24, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.6, 1.28, 10.6, 0.3, size=11, color=GRAY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.64), Inches(12.1), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()


def panel(slide, x, y, w, h, *, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def pill(slide, text, x, y, w, h, *, fill=MINT, color=TEAL, size=10):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    shape.line.width = Pt(0.7)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    return shape


def node(slide, text, x, y, w, h, *, fill=WHITE, line=TEAL, size=11):
    shape = panel(slide, x, y, w, h, fill=fill, line=line)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.07)
    tf.margin_right = Inches(0.07)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = NAVY
    return shape


def arrow(slide, x1, y1, x2, y2, *, color=TEAL, width=2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    try:
        line.line.end_arrowhead = True
    except Exception:
        pass
    return line


def add_picture_contain(slide, path: Path, x, y, w, h, *, border=True):
    if border:
        panel(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, fill=WHITE, line=LINE)
    if not path.exists():
        add_text(slide, "图片未找到", x, y + h / 2 - 0.15, w, 0.3, size=14, color=GRAY, align=PP_ALIGN.CENTER)
        return None
    with Image.open(path) as img:
        ratio = img.width / img.height
    box_ratio = w / h
    if ratio > box_ratio:
        draw_w = w
        draw_h = w / ratio
        draw_x = x
        draw_y = y + (h - draw_h) / 2
    else:
        draw_h = h
        draw_w = h * ratio
        draw_x = x + (w - draw_w) / 2
        draw_y = y
    return slide.shapes.add_picture(str(path), Inches(draw_x), Inches(draw_y), width=Inches(draw_w), height=Inches(draw_h))


def metric(slide, value, label, x, y, w=1.75, h=0.82, color=TEAL):
    panel(slide, x, y, w, h, fill=WHITE)
    add_text(slide, str(value), x + 0.15, y + 0.12, w - 0.3, 0.28, size=19, bold=True, color=color)
    add_text(slide, label, x + 0.15, y + 0.48, w - 0.3, 0.22, size=9, color=GRAY)


def make_deck():
    assets = prepare_assets()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 cover
    slide = prs.slides.add_slide(blank)
    set_bg(slide, RGBColor(235, 248, 255))
    add_text(slide, "多用户权限控制的\nRAG Agent 知识库系统", 0.72, 0.85, 7.5, 1.35, size=32, bold=True)
    add_text(slide, "面向企业文档的可溯源问答系统：从文档预处理开始提升 RAG 可靠性", 0.78, 2.44, 8.1, 0.35, size=15, color=GRAY)
    for i, label in enumerate(["文档预处理", "OCR fallback", "表格事实化", "权限检索", "防幻觉回答"]):
        pill(slide, label, 0.78 + i * 1.42, 3.13, 1.2, 0.34, fill=WHITE, color=[TEAL, BLUE, ORANGE, PINK, GREEN][i], size=9)
    add_text(slide, "Vue3 + FastAPI + ChromaDB + BGE + PaddleOCR + DeepSeek API", 0.8, 6.55, 8.2, 0.26, size=11, color=GRAY)
    # Cover flow
    y = 4.85
    xs = [7.8, 8.85, 9.9, 10.95, 12.0]
    labels = ["PDF", "OCR", "Chunk", "Rerank", "Answer"]
    colors = [TEAL, BLUE, ORANGE, PURPLE, GREEN]
    for i, (x, label, c) in enumerate(zip(xs, labels, colors)):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.62), Inches(0.62))
        shp.fill.solid()
        shp.fill.fore_color.rgb = c
        shp.line.fill.background()
        add_text(slide, label, x - 0.08, y + 0.72, 0.8, 0.22, size=8, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        if i < len(xs) - 1:
            arrow(slide, x + 0.63, y + 0.31, xs[i + 1] - 0.02, y + 0.31, color=GRAY, width=1.4)
    add_text(slide, "核心观点：RAG 的准确率上限，很大程度取决于文档预处理和切块质量。", 7.15, 1.25, 4.8, 0.78, size=22, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    # 2 pain points with real PDF cases
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "企业文档 RAG 的真实难点：不是“能不能问”，而是“能不能可靠地入库”", "扫描件、水印 PDF、表格密集标准文件会直接影响后续检索和回答质量。", "02 痛点")
    add_picture_contain(slide, assets["case_scan"], 0.75, 2.1, 3.25, 4.85)
    add_picture_contain(slide, assets["case_watermark"], 4.25, 2.1, 3.25, 4.85)
    pill(slide, "扫描件 PDF", 0.9, 2.22, 1.05, 0.32, fill=WHITE, color=BLUE)
    pill(slide, "带水印 PDF", 4.4, 2.22, 1.05, 0.32, fill=WHITE, color=PINK)
    panel(slide, 8.05, 2.1, 4.4, 4.85, fill=WHITE)
    add_text(slide, "直接切块会出什么问题？", 8.35, 2.38, 3.3, 0.3, size=19, bold=True)
    add_lines(
        slide,
        [
            "扫描件没有稳定文本层，普通 PDF 读取为空或乱码",
            "水印会干扰 OCR，表格线和正文容易混在一起",
            "标准文件的答案常在表格行中，简单按页切块会丢失结构",
            "如果入库片段质量差，后续 Embedding、Reranker、LLM 都难以补救",
        ],
        8.38,
        2.95,
        3.65,
        2.0,
        size=13,
        bullet=True,
    )
    add_text(slide, "因此本项目把预处理作为 RAG 链路的核心模块，而不是上传后的附属步骤。", 8.38, 5.7, 3.5, 0.48, size=15, bold=True, color=TEAL)

    # 3 architecture
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "系统架构：前端、后端、业务库、向量库与模型服务解耦", "用工程化结构承载文档上传、权限控制、检索问答和后续 Agent 扩展。", "03 架构")
    node(slide, "Vue3 前端\n上传 / 文档 / 问答 / 配置", 0.78, 2.45, 2.35, 1.0, fill=CYAN, line=BLUE)
    node(slide, "FastAPI 后端\n认证 / 文档 / 权限 / Chat", 4.0, 2.45, 2.35, 1.0, fill=WHITE, line=TEAL)
    node(slide, "SQLite\n用户 / 文档 / 权限 / 日志", 7.2, 1.9, 2.15, 0.9, fill=WHITE, line=ORANGE)
    node(slide, "ChromaDB\nchunk 向量 + metadata", 7.2, 3.35, 2.15, 0.9, fill=WHITE, line=PURPLE)
    node(slide, "模型服务\nBGE / Reranker / OCR / LLM", 10.18, 2.45, 2.25, 1.0, fill=MINT, line=GREEN)
    arrow(slide, 3.15, 2.95, 4.0, 2.95, color=BLUE)
    arrow(slide, 6.36, 2.75, 7.2, 2.35, color=ORANGE)
    arrow(slide, 6.36, 3.15, 7.2, 3.8, color=PURPLE)
    arrow(slide, 9.35, 2.35, 10.18, 2.75, color=GREEN)
    arrow(slide, 9.35, 3.8, 10.18, 3.15, color=GREEN)
    add_lines(
        slide,
        [
            "业务数据库记录用户、文档状态、权限、问答日志和审计日志",
            "向量库只负责 chunk 检索，并通过 metadata 做权限过滤",
            "模型能力封装为独立模块，便于替换本地模型或 API 服务",
        ],
        1.0,
        5.25,
        11.4,
        0.85,
        size=15,
        bullet=True,
    )

    # 4 end-to-end flow
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "完整链路：从上传文档到可溯源问答", "每个阶段都记录状态，便于排查“为什么检索不到、为什么回答不准”。", "04 流程")
    steps = [
        ("1 上传", "PDF / MD / TXT"),
        ("2 探测", "可读性 / 表格密度 / OCR需求"),
        ("3 解析", "PyMuPDF / pypdf / PaddleOCR"),
        ("4 增强", "表格事实化 / LLM归一化"),
        ("5 入库", "BGE Embedding / ChromaDB"),
        ("6 问答", "权限过滤 / Rerank / 引用"),
    ]
    for i, (title, desc) in enumerate(steps):
        x = 0.58 + i * 2.08
        node(slide, title, x, 2.35, 1.5, 0.56, fill=[CYAN, MINT, WHITE, WHITE, CYAN, MINT][i], line=[BLUE, TEAL, ORANGE, PINK, PURPLE, GREEN][i])
        add_text(slide, desc, x - 0.08, 3.08, 1.65, 0.42, size=10, color=GRAY, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arrow(slide, x + 1.52, 2.63, x + 2.02, 2.63, color=GRAY, width=1.4)
    panel(slide, 1.05, 4.4, 10.95, 1.45, fill=WHITE)
    add_text(slide, "可观测性贯穿全链路", 1.45, 4.7, 2.2, 0.3, size=18, bold=True, color=TEAL)
    add_text(slide, "文档状态、处理任务、chunk 数、质量分数、引用来源、问答日志、权限变更审计都可追踪。", 3.55, 4.72, 7.65, 0.3, size=15, color=INK)
    for i, text in enumerate(["status", "job", "quality", "source", "audit"]):
        pill(slide, text, 2.0 + i * 1.45, 5.28, 0.88, 0.28, fill=LIGHT, color=[TEAL, BLUE, ORANGE, PINK, GREEN][i], size=8)

    # 5 ingestion agent
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "轻量 Ingestion Agent：先探测，再选择处理策略", "当前版本以规则调度为主，后续可平滑迁移到 LangGraph 节点。", "05 Agent")
    node(slide, "文件上传", 0.9, 2.8, 1.4, 0.55, fill=CYAN, line=BLUE)
    node(slide, "轻量探测\n文本层 / 页数 / 表格密度", 3.0, 2.55, 2.0, 1.05, fill=MINT, line=TEAL)
    node(slide, "策略判断", 5.95, 2.7, 1.45, 0.75, fill=WHITE, line=ORANGE)
    node(slide, "普通 PDF\n文本解析", 8.2, 1.75, 1.55, 0.68, fill=WHITE, line=BLUE)
    node(slide, "扫描件\nOCR fallback", 8.2, 2.78, 1.55, 0.68, fill=WHITE, line=PINK)
    node(slide, "表格密集\n事实化切块", 8.2, 3.8, 1.55, 0.68, fill=WHITE, line=ORANGE)
    node(slide, "质量评分\nLLM归一化", 10.65, 2.78, 1.75, 0.68, fill=MINT, line=GREEN)
    arrow(slide, 2.3, 3.08, 3.0, 3.08, color=BLUE)
    arrow(slide, 5.0, 3.08, 5.95, 3.08, color=TEAL)
    arrow(slide, 7.4, 3.08, 8.2, 2.1, color=BLUE)
    arrow(slide, 7.4, 3.08, 8.2, 3.12, color=PINK)
    arrow(slide, 7.4, 3.08, 8.2, 4.15, color=ORANGE)
    arrow(slide, 9.75, 2.1, 10.65, 3.0, color=GREEN)
    arrow(slide, 9.75, 3.12, 10.65, 3.12, color=GREEN)
    arrow(slide, 9.75, 4.15, 10.65, 3.35, color=GREEN)
    add_lines(
        slide,
        [
            "避免所有文件都强制走 OCR 或 LLM，降低成本和等待时间",
            "对低质量 OCR、表格型切块进行结构化整理，提高检索可命中性",
            "处理状态回写任务表，前端能展示当前处于解析、OCR、LLM 还是 embedding",
        ],
        1.2,
        5.55,
        10.8,
        0.75,
        size=14,
        bullet=True,
    )

    # 6 table factization
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "表格事实化：把“看得懂的表格”转成“检索得到的知识事实”", "国标类文档的答案常以表格行形式存在，事实化能显著提升数值类问答稳定性。", "06 表格")
    panel(slide, 0.85, 2.15, 5.35, 3.95, fill=WHITE)
    add_text(slide, "普通读取后的片段", 1.15, 2.42, 2.1, 0.28, size=16, bold=True, color=RED)
    add_text(slide, "表7\n动物种类\n靶组织\n最大残留限量，μg/kg\n家禽\n蛋\n4\n2", 1.2, 3.0, 4.3, 1.8, size=20, color=INK)
    add_text(slide, "问题：行列关系被破坏，检索和生成都容易答错。", 1.2, 5.35, 4.5, 0.28, size=12, color=GRAY)
    arrow(slide, 6.35, 4.05, 7.0, 4.05, color=TEAL, width=2.2)
    panel(slide, 7.2, 2.15, 5.35, 3.95, fill=MINT, line=TEAL)
    add_text(slide, "事实化后的入库文本", 7.5, 2.42, 2.3, 0.28, size=16, bold=True, color=TEAL)
    add_text(
        slide,
        "来源文件：GB 31650.1-2022；页码：第6页；表号：表7。\n"
        "事实：青霉素在家禽蛋中的最大残留限量为 4 μg/kg。\n"
        "字段：兽药=青霉素；动物种类=家禽；靶组织=蛋；限量=4；单位=μg/kg。",
        7.55,
        3.05,
        4.35,
        1.5,
        size=15,
        color=NAVY,
    )
    add_text(slide, "优势：保留来源、页码、表号、字段关系，适合关键词、BM25 与向量检索共同命中。", 7.55, 5.35, 4.4, 0.4, size=12, color=GRAY)

    # 7 OCR/LLM progress UI
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "可视化入库进度：让重处理任务不再像“卡住”", "上传后展示解析、OCR、LLM 归一化、Embedding 等阶段，用户能判断等待是否正常。", "07 界面")
    add_picture_contain(slide, assets["ui_progress"], 0.75, 2.0, 7.1, 4.6)
    panel(slide, 8.25, 2.0, 4.0, 4.6, fill=WHITE)
    add_text(slide, "界面展示的关键信息", 8.55, 2.3, 2.4, 0.3, size=18, bold=True)
    add_lines(
        slide,
        [
            "是否开启“大模型整理 OCR / 表格切块”",
            "当前正在处理哪个文档",
            "当前阶段：OCR / LLM 归一化 / 入库",
            "处理进度百分比，降低用户等待焦虑",
        ],
        8.55,
        2.95,
        3.1,
        1.8,
        size=13,
        bullet=True,
    )
    add_text(slide, "适合活动现场演示：评委能直接看到文档入库不是黑盒。", 8.55, 5.65, 3.2, 0.42, size=14, bold=True, color=TEAL)

    # 8 upload UI
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "系统界面：文档入库工作台", "上传、状态统计、模型配置、最近入库文档在同一工作台中集中展示。", "08 界面")
    add_picture_contain(slide, assets["ui_upload"], 0.72, 2.0, 8.65, 4.75)
    panel(slide, 9.65, 2.0, 2.75, 4.75, fill=WHITE)
    add_text(slide, "演示讲法", 9.92, 2.28, 1.4, 0.3, size=18, bold=True, color=TEAL)
    add_lines(
        slide,
        [
            "登录后上传 PDF",
            "系统自动解析并入库",
            "文档列表显示状态、chunk 数和上传者",
            "管理员可删除文档或查看切块",
        ],
        9.9,
        2.9,
        2.05,
        1.8,
        size=13,
        bullet=True,
    )
    metric(slide, "BGE", "Embedding", 9.95, 5.35, 1.0, 0.62, color=BLUE)
    metric(slide, "OCR", "Fallback", 11.1, 5.35, 1.0, 0.62, color=PINK)

    # 9 chat UI
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "系统界面：问答测试与证据核查", "对话式问答支持选择任意数量文档，回答按“分类判断、结论、依据、引用”呈现。", "09 界面")
    add_picture_contain(slide, assets["ui_chat"], 0.72, 2.0, 8.8, 4.8)
    panel(slide, 9.82, 2.0, 2.65, 4.8, fill=WHITE)
    add_text(slide, "防幻觉设计", 10.1, 2.3, 1.7, 0.3, size=18, bold=True, color=PINK)
    add_lines(
        slide,
        [
            "先检索权限内文档",
            "证据不足则回答“未找到”",
            "返回引用来源，便于人工复核",
            "检索依据默认收起，界面不干扰主答案",
        ],
        10.05,
        2.95,
        2.05,
        1.9,
        size=13,
        bullet=True,
    )
    add_text(slide, "适合企业内部文档：宁可不答，也不编造。", 10.05, 5.75, 2.05, 0.38, size=13, bold=True, color=TEAL)

    # 10 retrieval
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "检索增强：向量检索解决语义，BM25 兜住标准号和数值", "食品国标问答中，添加剂名称、标准号、限量值对关键词召回更敏感。", "10 检索")
    node(slide, "用户问题", 0.8, 3.0, 1.4, 0.6, fill=CYAN, line=BLUE)
    node(slide, "Query 归一化\ntokenizer", 2.9, 2.05, 1.7, 0.75, fill=WHITE, line=TEAL)
    node(slide, "向量召回\nBGE + ChromaDB", 5.25, 1.55, 1.8, 0.75, fill=WHITE, line=PURPLE)
    node(slide, "BM25 / 关键词\n标准号、术语、数值", 5.25, 3.0, 1.8, 0.8, fill=WHITE, line=ORANGE)
    node(slide, "候选融合\nscore merge", 7.85, 2.25, 1.6, 0.75, fill=MINT, line=TEAL)
    node(slide, "BGE Reranker\n精排 top-k", 10.05, 2.25, 1.65, 0.75, fill=WHITE, line=PINK)
    node(slide, "LLM 生成\n引用来源", 10.05, 4.0, 1.65, 0.75, fill=MINT, line=GREEN)
    arrow(slide, 2.2, 3.3, 2.9, 2.42, color=BLUE)
    arrow(slide, 4.6, 2.42, 5.25, 1.92, color=PURPLE)
    arrow(slide, 4.6, 2.42, 5.25, 3.4, color=ORANGE)
    arrow(slide, 7.05, 1.92, 7.85, 2.52, color=PURPLE)
    arrow(slide, 7.05, 3.4, 7.85, 2.68, color=ORANGE)
    arrow(slide, 9.45, 2.63, 10.05, 2.63, color=PINK)
    arrow(slide, 10.88, 3.0, 10.88, 4.0, color=GREEN)
    add_text(slide, "核心策略：不依赖单一路径召回，先扩大候选，再精排压缩上下文。", 2.1, 5.8, 8.8, 0.38, size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    # 11 permissions
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "权限系统：不只前端隐藏，检索阶段也必须过滤", "用户只能在自己、部门或显式授权范围内召回文档。", "11 权限")
    node(slide, "用户登录\nJWT", 0.9, 2.55, 1.55, 0.75, fill=CYAN, line=BLUE)
    node(slide, "身份解析\nadmin / user", 3.0, 2.55, 1.65, 0.75, fill=WHITE, line=TEAL)
    node(slide, "可访问文档集合\nowner / user / group", 5.35, 2.35, 2.05, 1.05, fill=MINT, line=TEAL)
    node(slide, "Chroma metadata filter\ndocument_id in [...]", 8.1, 2.35, 2.1, 1.05, fill=WHITE, line=PURPLE)
    node(slide, "RAG 回答\n仅基于权限内证据", 10.85, 2.55, 1.65, 0.75, fill=WHITE, line=GREEN)
    arrow(slide, 2.45, 2.92, 3.0, 2.92, color=BLUE)
    arrow(slide, 4.65, 2.92, 5.35, 2.92, color=TEAL)
    arrow(slide, 7.4, 2.92, 8.1, 2.92, color=PURPLE)
    arrow(slide, 10.2, 2.92, 10.85, 2.92, color=GREEN)
    panel(slide, 1.15, 4.85, 10.8, 0.92, fill=WHITE)
    add_text(slide, "审计日志", 1.45, 5.12, 1.0, 0.25, size=15, bold=True, color=PINK)
    add_text(slide, "记录权限授予、删除、用户组成员变化，后续可用于企业知识库合规审查。", 2.5, 5.12, 7.7, 0.25, size=14, color=INK)

    # 12 quality report
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "文档处理质量报告：把入库质量量化出来", "不是只看“上传成功”，而是看 chunk 是否可检索、可引用、可维护。", "12 质量")
    labels = ["文本可读性", "表格结构", "OCR 噪声", "来源完整", "可检索性"]
    vals = [86, 78, 72, 94, 88]
    for i, (label, val) in enumerate(zip(labels, vals)):
        y = 2.25 + i * 0.78
        add_text(slide, label, 1.1, y + 0.06, 1.2, 0.2, size=12, color=INK, bold=True)
        bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.6), Inches(y), Inches(6.8), Inches(0.28))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(226, 232, 240)
        bg.line.fill.background()
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.6), Inches(y), Inches(6.8 * val / 100), Inches(0.28))
        bar.fill.solid()
        bar.fill.fore_color.rgb = [TEAL, BLUE, ORANGE, GREEN, PURPLE][i]
        bar.line.fill.background()
        add_text(slide, f"{val}", 9.65, y - 0.02, 0.5, 0.25, size=12, color=NAVY, bold=True)
    panel(slide, 10.45, 2.25, 1.8, 3.4, fill=MINT, line=TEAL)
    add_text(slide, "质量报告用于", 10.7, 2.55, 1.2, 0.28, size=15, bold=True, color=TEAL)
    add_lines(slide, ["发现 OCR 噪声", "定位低质切块", "决定是否重跑 LLM 整理", "支撑评测复盘"], 10.68, 3.1, 1.25, 1.6, size=11, bullet=True)

    # 13 results
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "阶段性效果：预处理增强后，RAG 链路达到可接受演示精度", "评测结果主要用于迭代方向判断，不包装成最终生产指标。", "13 效果")
    # custom bar chart
    bars = [
        ("初始链路", 71.4, GRAY),
        ("BM25 调优", 80.0, BLUE),
        ("v1 修正", 90.0, TEAL),
        ("预处理增强", 95.7, GREEN),
        ("v2 扩展集", 90.4, PURPLE),
    ]
    base_y = 5.85
    for i, (label, val, c) in enumerate(bars):
        x = 1.0 + i * 2.2
        h = 3.0 * val / 100
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(base_y - h), Inches(1.05), Inches(h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = c
        rect.line.fill.background()
        add_text(slide, f"{val:.1f}%", x - 0.08, base_y - h - 0.36, 1.25, 0.22, size=13, bold=True, color=c, align=PP_ALIGN.CENTER)
        add_text(slide, label, x - 0.35, base_y + 0.16, 1.75, 0.3, size=10, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, "提升来源：表格事实化、OCR fallback、LLM 归一化、BM25 混合召回、Reranker 精排、防幻觉提示词。", 1.0, 6.65, 11.1, 0.28, size=13, color=GRAY, align=PP_ALIGN.CENTER)

    # 14 highlights
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "项目亮点：从 RAG Demo 到工程化知识库系统", "展示的不只是模型调用，而是一条可解释、可维护、可扩展的应用链路。", "14 亮点")
    highlights = [
        ("文档预处理优先", "扫描件、水印、表格密集 PDF 均有处理策略"),
        ("可溯源 RAG", "答案返回引用来源，证据不足时拒绝编造"),
        ("权限检索", "后端接口和向量检索阶段共同生效"),
        ("混合检索", "向量 + BM25 + 关键词 + Reranker"),
        ("可观测入库", "任务表、质量分数、进度条、错误日志"),
        ("Agent 化路线", "当前轻量调度，后续 LangGraph 节点化"),
    ]
    for i, (head, desc) in enumerate(highlights):
        x = 0.9 + (i % 3) * 4.0
        y = 2.18 + (i // 3) * 2.1
        panel(slide, x, y, 3.35, 1.42, fill=WHITE, line=[TEAL, BLUE, PINK, ORANGE, GREEN, PURPLE][i])
        add_text(slide, head, x + 0.25, y + 0.26, 2.45, 0.25, size=16, bold=True, color=[TEAL, BLUE, PINK, ORANGE, GREEN, PURPLE][i])
        add_text(slide, desc, x + 0.25, y + 0.73, 2.65, 0.36, size=12, color=GRAY)

    # 15 roadmap
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "后续规划：把轻量调度升级为真正的 LangGraph Agent", "在当前可演示链路上继续增强自动决策、工具调用和评测闭环。", "15 规划")
    roadmap = [
        ("短期", "完善入库队列与 GPU 并发控制\n沉淀更大规模评测集"),
        ("中期", "LangGraph 化：探测、OCR、切块、检索、生成拆成节点\n支持失败重试和可视化轨迹"),
        ("长期", "接入企业 SSO / PostgreSQL / 对象存储\n支持多知识库、多租户和审计报表"),
    ]
    for i, (head, body) in enumerate(roadmap):
        x = 1.0 + i * 4.05
        panel(slide, x, 2.35, 3.25, 2.55, fill=WHITE, line=[TEAL, BLUE, PURPLE][i])
        pill(slide, head, x + 0.28, 2.67, 0.86, 0.34, fill=MINT if i == 0 else CYAN, color=[TEAL, BLUE, PURPLE][i])
        add_text(slide, body, x + 0.35, 3.35, 2.45, 0.8, size=14, color=INK)
    add_text(slide, "汇报收束句：这不是一个只会“接大模型接口”的 Demo，而是围绕企业文档质量、权限和可溯源构建的 RAG 工程化原型。", 1.05, 6.25, 11.3, 0.38, size=17, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    DOCS.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    out = make_deck()
    print(out)
