from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "RAG_Agent_KB_System_活动汇报_初版.pptx"
SCAN_CASE_IMAGE = Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/fd132a56-daf2-42a7-80b7-bd9b1952ca74.png")
WATERMARK_CASE_IMAGE = Path(r"E:/xwechat_files/wxid_2g2jkeczgtwr22_bbcf/temp/InputTemp/80fb6e19-cc2c-4c0b-8475-578db475a3f8.png")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(18, 30, 54)
BLUE = RGBColor(29, 111, 214)
TEAL = RGBColor(15, 118, 110)
MINT = RGBColor(229, 247, 244)
CYAN = RGBColor(228, 242, 255)
PINK = RGBColor(224, 22, 91)
ORANGE = RGBColor(245, 158, 11)
GRAY = RGBColor(92, 108, 130)
LIGHT = RGBColor(246, 249, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(210, 222, 235)
GREEN = RGBColor(22, 163, 74)
RED = RGBColor(220, 38, 38)

FONT = "Microsoft YaHei"


def set_bg(slide, color=LIGHT):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 20,
    bold: bool = False,
    color= NAVY,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_multiline(slide, lines, x, y, w, h, *, size=16, color=NAVY, bullet=False, gap=0.12):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(5)
        p.level = 0
        if bullet:
            p.text = f"• {line}"
    return box


def add_title(slide, title: str, subtitle: str | None = None, section: str | None = None):
    if section:
        pill(slide, section, 0.58, 0.34, 1.45, 0.34, fill=CYAN, color=BLUE, size=10)
    add_text(slide, title, 0.58, 0.72, 8.8, 0.55, size=24, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, subtitle, 0.6, 1.26, 9.6, 0.34, size=11, color=GRAY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.62), Inches(11.85), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()


def pill(slide, text, x, y, w, h, *, fill=MINT, color=TEAL, size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    shape.line.width = Pt(0.8)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    return shape


def panel(slide, x, y, w, h, *, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def metric(slide, value, label, x, y, w=1.85, h=0.9, color=TEAL):
    panel(slide, x, y, w, h, fill=WHITE)
    add_text(slide, str(value), x + 0.16, y + 0.14, w - 0.3, 0.32, size=20, bold=True, color=color)
    add_text(slide, label, x + 0.16, y + 0.52, w - 0.28, 0.24, size=9, color=GRAY)


def arrow(slide, x1, y1, x2, y2, color=TEAL):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    try:
        line.line.end_arrowhead = True
    except Exception:
        pass
    return line


def flow_node(slide, text, x, y, w, h, *, fill=WHITE, color=TEAL, size=11):
    shape = panel(slide, x, y, w, h, fill=fill, line=color)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = NAVY
    return shape


def add_case_image(slide, path: Path, x: float, y: float, w: float, h: float):
    panel(slide, x - 0.04, y - 0.04, w + 0.08, h + 0.08, fill=WHITE, line=LINE)
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    else:
        add_text(slide, "截图文件未找到\n请替换为系统案例截图", x + 0.25, y + h / 2 - 0.25, w - 0.5, 0.5, size=13, color=GRAY, align=PP_ALIGN.CENTER)


def make_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 cover
    slide = prs.slides.add_slide(blank)
    set_bg(slide, RGBColor(236, 247, 255))
    add_text(slide, "多用户权限控制的\nRAG Agent 知识库系统", 0.72, 1.0, 7.3, 1.35, size=34, bold=True, color=NAVY)
    add_text(slide, "面向企业文档的预处理增强、权限检索与可溯源问答", 0.78, 2.55, 7.6, 0.42, size=16, color=GRAY)
    pill(slide, "AIGC 应用项目", 0.78, 3.18, 1.55, 0.38, fill=WHITE, color=BLUE)
    pill(slide, "文档预处理增强", 2.52, 3.18, 1.85, 0.38, fill=WHITE, color=TEAL)
    pill(slide, "企业知识库 RAG", 4.56, 3.18, 1.85, 0.38, fill=WHITE, color=PINK)
    add_text(slide, "Vue3 · FastAPI · ChromaDB · BGE · PaddleOCR · DeepSeek API", 0.8, 6.55, 8.5, 0.28, size=11, color=GRAY)
    # visual circuit
    for i, (x, y, c) in enumerate([(9.0, 0.9, TEAL), (10.45, 1.55, BLUE), (9.55, 2.55, PINK), (11.2, 3.15, ORANGE), (10.0, 4.35, GREEN)]):
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.62), Inches(0.62))
        node.fill.solid(); node.fill.fore_color.rgb = c; node.line.fill.background()
        add_text(slide, str(i + 1), x, y + 0.13, 0.62, 0.22, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    arrow(slide, 9.62, 1.2, 10.45, 1.86, BLUE)
    arrow(slide, 10.75, 2.12, 9.85, 2.55, TEAL)
    arrow(slide, 10.05, 2.85, 11.2, 3.46, PINK)
    arrow(slide, 11.2, 3.78, 10.3, 4.35, ORANGE)
    add_text(slide, "从“能问”走向“能准、可查、可控”", 8.52, 5.45, 3.8, 0.6, size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    # 2 background
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "项目背景：企业文档 RAG 的难点在入库质量", "国标、扫描件、表格型 PDF 中，错误往往发生在检索之前。", "02 背景")
    add_text(slide, "普通 RAG 常见问题", 0.8, 2.0, 3.0, 0.35, size=18, bold=True, color=PINK)
    add_multiline(slide, ["PDF 文本层乱码或不可读", "表格按行错乱，对象和值无法对应", "OCR 断行、错字、单位/上标识别错误", "缺少页码、表号和来源信息", "证据不足时模型容易编造答案"], 0.82, 2.52, 4.6, 2.2, size=15, color=NAVY, bullet=True)
    panel(slide, 6.2, 2.0, 5.8, 3.1, fill=WHITE)
    add_text(slide, "核心判断", 6.55, 2.28, 1.6, 0.26, size=12, bold=True, color=TEAL)
    add_text(slide, "RAG 的准确率上限，很大程度取决于文档预处理和切块质量。", 6.55, 2.75, 4.9, 0.86, size=24, bold=True, color=NAVY)
    add_text(slide, "因此本项目优先优化“文档入库前”的 OCR、表格、切块质量与可溯源 metadata。", 6.58, 4.08, 4.75, 0.55, size=13, color=GRAY)
    metric(slide, "PDF", "复杂格式", 0.8, 5.7)
    metric(slide, "OCR", "扫描件 fallback", 2.9, 5.7, color=BLUE)
    metric(slide, "表格", "事实化转写", 5.0, 5.7, color=ORANGE)
    metric(slide, "引用", "页码可溯源", 7.1, 5.7, color=PINK)
    metric(slide, "拒答", "降低幻觉", 9.2, 5.7, color=GREEN)

    # 3 goals
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "项目目标：从本地 RAG Demo 升级为工程化知识库系统", "目标是展示企业级大模型应用开发能力，而不是只完成一次问答。", "03 目标")
    cols = [
        ("文档处理", ["文件上传与类型识别", "PDF / OCR / 表格解析", "切块质量评分", "LLM 切块整理"]),
        ("权限检索", ["JWT 登录鉴权", "用户 / 部门权限", "后端接口过滤", "检索阶段 document_id 过滤"]),
        ("可溯源问答", ["混合检索 + Reranker", "分类判断 / 结论 / 依据 / 引用", "证据不足返回未找到", "问答日志与审计日志"]),
    ]
    for idx, (head, items) in enumerate(cols):
        x = 0.78 + idx * 4.12
        panel(slide, x, 2.0, 3.65, 3.8, fill=WHITE)
        pill(slide, head, x + 0.28, 2.28, 1.35, 0.36, fill=MINT if idx != 1 else CYAN, color=[TEAL, BLUE, PINK][idx])
        add_multiline(slide, items, x + 0.32, 2.96, 3.05, 1.9, size=13, color=NAVY, bullet=True)
    add_text(slide, "一句话定位：带权限、可观测、可评测、可扩展 Agent 的企业知识库系统。", 1.05, 6.35, 11.0, 0.36, size=19, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    # 4 architecture
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "系统架构：前端、后端、业务库、向量库与模型服务解耦", "模块化分层便于展示、测试和后续替换模型或数据库。", "04 架构")
    flow_node(slide, "Vue3 前端\n上传 / 列表 / 问答 / 配置", 0.8, 2.25, 2.25, 1.0, fill=WHITE, color=BLUE)
    flow_node(slide, "FastAPI 后端\nAPI / Service / Schema", 3.55, 2.25, 2.25, 1.0, fill=WHITE, color=TEAL)
    flow_node(slide, "RAG 模块\nLoader / Chunker / Retriever", 6.3, 2.25, 2.25, 1.0, fill=WHITE, color=PINK)
    flow_node(slide, "模型服务\nBGE / OCR / DeepSeek", 9.05, 2.25, 2.25, 1.0, fill=WHITE, color=ORANGE)
    arrow(slide, 3.05, 2.75, 3.55, 2.75, BLUE)
    arrow(slide, 5.8, 2.75, 6.3, 2.75, TEAL)
    arrow(slide, 8.55, 2.75, 9.05, 2.75, PINK)
    flow_node(slide, "SQLite\n用户 / 权限 / 任务 / 日志", 2.0, 4.35, 2.65, 1.0, fill=CYAN, color=BLUE)
    flow_node(slide, "ChromaDB\n向量 / metadata 检索", 5.3, 4.35, 2.65, 1.0, fill=MINT, color=TEAL)
    flow_node(slide, "data/uploads\n原始文件本地存储", 8.6, 4.35, 2.65, 1.0, fill=WHITE, color=GRAY)
    arrow(slide, 4.6, 3.25, 3.3, 4.35, GRAY)
    arrow(slide, 7.42, 3.25, 6.62, 4.35, GRAY)
    arrow(slide, 10.15, 3.25, 9.92, 4.35, GRAY)
    add_text(slide, "特点：配置集中在 .env / config.py；核心逻辑不堆在 main.py；权限在后端和检索阶段同时生效。", 0.95, 6.28, 11.2, 0.36, size=14, color=GRAY, align=PP_ALIGN.CENTER)

    # 5 full flow
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "完整链路：从上传文档到可溯源问答", "重点高亮的不是最后生成答案，而是答案之前的可控证据链。", "05 流程")
    steps = [
        "登录\nJWT",
        "上传\nPDF/TXT/MD",
        "Agent\n文件探测",
        "解析/OCR\n表格抽取",
        "切块质量\n评分",
        "Embedding\n入库",
        "权限过滤\n检索",
        "混合召回\nReranker",
        "证据约束\n生成答案",
    ]
    x0 = 0.55
    for i, step in enumerate(steps):
        x = x0 + (i % 5) * 2.48
        y = 2.05 + (i // 5) * 1.75
        color = TEAL if 2 <= i <= 4 else BLUE if i < 2 else PINK
        flow_node(slide, step, x, y, 1.72, 0.76, fill=WHITE, color=color, size=10)
        if i not in (4, 8):
            nx = x0 + ((i + 1) % 5) * 2.48
            ny = 2.05 + ((i + 1) // 5) * 1.75
            if i == 4:
                continue
            arrow(slide, x + 1.72, y + 0.38, nx, ny + 0.38, GRAY)
    arrow(slide, x0 + 4 * 2.48 + 0.86, 2.82, x0 + 4 * 2.48 + 0.86, 3.8, GRAY)
    arrow(slide, x0 + 4 * 2.48, 4.18, x0 + 3 * 2.48 + 1.72, 4.18, GRAY)
    arrow(slide, x0 + 3 * 2.48, 4.18, x0 + 2 * 2.48 + 1.72, 4.18, GRAY)
    arrow(slide, x0 + 2 * 2.48, 4.18, x0 + 1 * 2.48 + 1.72, 4.18, GRAY)
    add_text(slide, "预处理增强区", 5.45, 1.58, 2.0, 0.28, size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    panel(slide, 5.1, 1.88, 5.35, 1.12, fill=RGBColor(240, 253, 250), line=TEAL)
    add_text(slide, "这一区域决定了后续检索能否找到正确证据。", 5.45, 5.45, 5.2, 0.32, size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    # 6 preprocessing strategy
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "重点一：入库前先探测，再选择处理策略", "轻量 Ingestion Agent 不盲目调用 OCR 或 LLM，而是根据文件质量决策。", "06 预处理")
    left = ["文件类型识别", "PDF 文本层可读性判断", "表格密度检测", "OCR 需求判断", "LLM 归一化需求判断", "切块质量评分与风险提示"]
    add_text(slide, "探测信号", 0.8, 2.0, 2.0, 0.3, size=18, bold=True, color=TEAL)
    add_multiline(slide, left, 0.85, 2.5, 4.3, 2.7, size=15, color=NAVY, bullet=True)
    strategies = [
        ("可读 PDF", "PyMuPDF 提取正文、页码和表格"),
        ("文本层异常", "PaddleOCR fallback，保留页级来源"),
        ("表格密集", "表格行转成可检索事实"),
        ("低质量切块", "LLM 结构化归一化，可通过上传开关控制"),
    ]
    add_text(slide, "处理策略", 6.0, 2.0, 2.0, 0.3, size=18, bold=True, color=BLUE)
    for idx, (a, b) in enumerate(strategies):
        y = 2.45 + idx * 0.86
        pill(slide, a, 6.05, y, 1.25, 0.32, fill=CYAN, color=BLUE, size=10)
        add_text(slide, b, 7.55, y + 0.02, 4.3, 0.3, size=13, color=NAVY)
    add_text(slide, "结果：每个 chunk 都尽量带上文件名、页码、表格信息、质量分数和处理来源。", 1.0, 6.1, 10.9, 0.42, size=20, bold=True, color=PINK, align=PP_ALIGN.CENTER)

    # 7 table/OCR before after
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "重点二：表格不是简单抽文字，而是转成可检索事实", "国标类文档的答案经常藏在表格中，必须保留对象、限量值和来源。", "07 表格增强")
    panel(slide, 0.7, 2.05, 5.55, 3.55, fill=WHITE)
    add_text(slide, "处理前：表格行被拆散", 1.0, 2.28, 3.0, 0.28, size=16, bold=True, color=PINK)
    add_multiline(slide, ["胶基糖果", "己二酸", "4.0 g/kg", "页码和表格关系弱", "检索时容易只命中数值或名称"], 1.05, 2.88, 4.6, 1.85, size=15, color=NAVY, bullet=True)
    panel(slide, 7.0, 2.05, 5.55, 3.55, fill=RGBColor(240, 253, 250), line=TEAL)
    add_text(slide, "处理后：事实化 + 可溯源", 7.32, 2.28, 3.0, 0.28, size=16, bold=True, color=TEAL)
    add_text(slide, "己二酸在胶基糖果中的最大使用量为 4.0 g/kg。\n来源文件：GB 2760-2024；页码：第 X 页；表格：表 X。", 7.35, 2.95, 4.55, 1.05, size=17, bold=True, color=NAVY)
    add_multiline(slide, ["对象和值绑定", "引用页码清晰", "更适合 BM25 / 关键词 / 向量共同召回"], 7.35, 4.35, 4.55, 0.88, size=13, color=GRAY, bullet=True)
    arrow(slide, 6.28, 3.82, 6.95, 3.82, TEAL)

    # 8 LLM normalization and quality
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "重点三：OCR 与 LLM 切块整理提高低质量文档可用性", "对扫描件和复杂表格，预处理质量直接决定后续答案质量。", "08 OCR/LLM")
    items = [
        ("OCR fallback", "文本层不可读、扫描件或乱码 PDF 自动进入 OCR 路径"),
        ("断行与单位修复", "处理奇怪换行、μg/kg、10 的上标等关键形式"),
        ("质量评分", "给 chunk 标记 high / medium / low 和风险 flags"),
        ("LLM 归一化", "只对候选低质量 OCR / 表格块整理，避免所有块强制调用"),
        ("进度可视化", "前端显示 OCR、LLM、Embedding、入库阶段和当前进度"),
    ]
    for idx, (head, body) in enumerate(items):
        y = 2.0 + idx * 0.74
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(y), Inches(0.34), Inches(0.34))
        shape.fill.solid(); shape.fill.fore_color.rgb = [TEAL, BLUE, ORANGE, PINK, GREEN][idx]; shape.line.fill.background()
        add_text(slide, head, 1.45, y - 0.01, 2.0, 0.3, size=15, bold=True, color=NAVY)
        add_text(slide, body, 3.48, y + 0.01, 7.8, 0.28, size=13, color=GRAY)
    panel(slide, 8.9, 5.25, 3.0, 0.88, fill=WHITE)
    add_text(slide, "设计原则", 9.15, 5.43, 1.0, 0.22, size=11, bold=True, color=TEAL)
    add_text(slide, "把算力用在最影响准确率的块上", 9.15, 5.75, 2.35, 0.22, size=12, bold=True, color=NAVY)

    # 9 quality report
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "质量报告：让文档入库从黑盒变成可观察流程", "每个文档不只记录“成功/失败”，还输出可解释的处理质量。", "09 可观测")
    metric(slide, "健康分", "综合质量评分", 0.85, 2.05, 2.0, 1.0, color=TEAL)
    metric(slide, "风险等级", "低 / 中 / 高", 3.05, 2.05, 2.0, 1.0, color=PINK)
    metric(slide, "OCR 块", "扫描件处理占比", 5.25, 2.05, 2.0, 1.0, color=BLUE)
    metric(slide, "表格块", "结构化处理重点", 7.45, 2.05, 2.0, 1.0, color=ORANGE)
    metric(slide, "LLM", "归一化块数", 9.65, 2.05, 2.0, 1.0, color=GREEN)
    stages = ["Agent 探测", "解析 / OCR", "表格归一化", "分块与入库"]
    for idx, stage in enumerate(stages):
        flow_node(slide, stage, 1.1 + idx * 2.9, 4.15, 2.2, 0.68, fill=WHITE, color=TEAL if idx < 3 else BLUE)
        if idx < len(stages) - 1:
            arrow(slide, 3.3 + idx * 2.9, 4.49, 4.0 + idx * 2.9, 4.49, GRAY)
    add_text(slide, "用途：快速定位失败原因、判断是否需要重建向量、筛选需要人工抽查的高风险文档。", 1.0, 6.0, 11.0, 0.35, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    # 10 real case examples
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "真实案例：扫描件与水印 PDF 是预处理必须解决的问题", "这类文档如果直接粗切块，表格关系、页码来源和关键数值都可能丢失。", "10 案例")
    add_case_image(slide, SCAN_CASE_IMAGE, 0.78, 1.95, 3.85, 4.65)
    add_case_image(slide, WATERMARK_CASE_IMAGE, 5.0, 1.95, 3.85, 4.65)
    add_text(slide, "PDF 扫描件", 0.95, 6.75, 3.4, 0.26, size=12, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_text(slide, "带水印 PDF", 5.18, 6.75, 3.4, 0.26, size=12, bold=True, color=PINK, align=PP_ALIGN.CENTER)
    panel(slide, 9.25, 2.1, 3.25, 3.9, fill=WHITE)
    add_text(slide, "处理目标", 9.55, 2.42, 1.3, 0.28, size=16, bold=True, color=NAVY)
    add_multiline(
        slide,
        [
            "识别是否需要 OCR fallback",
            "去除或弱化水印干扰",
            "把表格行转成事实文本",
            "保留页码、表号、文件名",
            "对低质量块触发 LLM 归一化",
        ],
        9.55,
        2.9,
        2.55,
        1.85,
        size=12,
        color=NAVY,
        bullet=True,
    )
    add_text(slide, "核心价值：让“看得到但机器读不准”的文档，也能变成可检索、可引用的知识。", 9.45, 5.35, 2.65, 0.42, size=13, bold=True, color=TEAL)

    # 11 retrieval anti hallucination
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "检索增强与防幻觉：让答案有依据", "标准号、限量值、单位和专业术语需要比普通语义检索更强的召回策略。", "11 检索")
    nodes = [("用户问题", BLUE), ("向量召回", TEAL), ("BM25/关键词", ORANGE), ("Reranker", PINK), ("证据判断", GREEN), ("结构化答案", BLUE)]
    for idx, (text, color) in enumerate(nodes):
        x = 0.75 + idx * 2.05
        flow_node(slide, text, x, 2.25, 1.48, 0.72, fill=WHITE, color=color, size=10)
        if idx < len(nodes) - 1:
            arrow(slide, x + 1.48, 2.61, x + 2.05, 2.61, GRAY)
    add_text(slide, "防幻觉策略", 0.9, 4.02, 2.0, 0.3, size=18, bold=True, color=PINK)
    add_multiline(slide, ["证据不足时返回“未找到”", "答案必须包含分类判断、结论、依据、引用", "检索依据默认折叠，但可展开核查", "问答日志记录问题、答案和引用来源"], 0.95, 4.48, 5.25, 1.45, size=14, color=NAVY, bullet=True)
    add_text(slide, "适配场景", 7.05, 4.02, 2.0, 0.3, size=18, bold=True, color=TEAL)
    add_multiline(slide, ["GB 标准号查询", "添加剂/污染物/兽药限量", "跨文档引用问题", "企业内部制度问答"], 7.1, 4.48, 4.4, 1.45, size=14, color=NAVY, bullet=True)

    # 12 evaluation
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "阶段性效果：预处理增强后，RAG 链路达到可接受准确率", "评测用于发现问题，不是简单追求单次分数。", "12 效果")
    metric(slide, "500", "v2 测试集问题数", 0.9, 2.1, 2.1, 1.05, color=BLUE)
    metric(slide, "90.40%", "清理后离线重判", 3.35, 2.1, 2.1, 1.05, color=TEAL)
    metric(slide, "6类", "覆盖问题类型", 5.8, 2.1, 2.1, 1.05, color=PINK)
    metric(slide, "95%+", "小规模增强实验水平", 8.25, 2.1, 2.2, 1.05, color=GREEN)
    add_text(slide, "覆盖问题", 0.95, 4.0, 1.5, 0.3, size=16, bold=True, color=NAVY)
    add_multiline(slide, ["添加剂使用量", "污染物限量", "兽药残留", "产品标准", "否定问题", "跨文档问题"], 0.98, 4.45, 4.0, 1.25, size=13, color=NAVY, bullet=True)
    add_text(slide, "剩余误差来源", 6.0, 4.0, 2.0, 0.3, size=16, bold=True, color=NAVY)
    add_multiline(slide, ["复杂长表格", "跨文档链式检索", "OCR 质量波动", "题目本身歧义或脏数据"], 6.02, 4.45, 4.8, 1.25, size=13, color=NAVY, bullet=True)
    add_text(slide, "结论：当前准确率已经可用于项目展示，后续重点转向 Agent 化和工程稳定性。", 1.0, 6.35, 11.1, 0.34, size=17, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    # 13 frontend demo
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "可视化工作台：让评委能快速看懂系统能力", "前端按工作区拆分，避免所有功能挤在同一页。", "13 前端")
    modules = [
        ("文档上传", "LLM 整理开关、上传进度、OCR/Embedding 状态"),
        ("文档列表", "上传者、状态、分块数、质量报告、管理员删除"),
        ("问答测试", "多文档选择、对话式问答、四行结构化答案"),
        ("系统配置", "能力矩阵、流程看板、模型配置、权限与审计"),
    ]
    for idx, (head, body) in enumerate(modules):
        x = 0.85 + (idx % 2) * 5.75
        y = 2.05 + (idx // 2) * 1.75
        panel(slide, x, y, 5.0, 1.22, fill=WHITE)
        pill(slide, head, x + 0.28, y + 0.24, 1.25, 0.32, fill=CYAN if idx % 2 else MINT, color=BLUE if idx % 2 else TEAL, size=10)
        add_text(slide, body, x + 0.32, y + 0.72, 4.2, 0.28, size=13, color=NAVY)
    add_text(slide, "建议现场演示：先打开系统配置页讲全局，再进入文档列表看质量报告，最后做一次问答。", 1.0, 6.28, 11.2, 0.35, size=16, bold=True, color=GRAY, align=PP_ALIGN.CENTER)

    # 14 highlights
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "项目亮点：从 RAG Demo 到工程化知识库系统", "亮点不只在模型调用，而在完整链路、权限和可观测性。", "14 亮点")
    highlights = [
        ("文档预处理增强", "文件探测、OCR fallback、表格事实化和 LLM 切块整理。"),
        ("国标文档适配", "针对标准号、限量值、单位、表格页码设计检索策略。"),
        ("权限感知 RAG", "用户只能检索自己有权限的文档，后端和检索层同时过滤。"),
        ("混合检索重排", "融合 BGE 向量、BM25、关键词和 Reranker。"),
        ("可观测可评测", "任务表、质量报告、问答日志、审计日志和 500 条测试集。"),
    ]
    for idx, (head, body) in enumerate(highlights):
        y = 2.0 + idx * 0.74
        add_text(slide, f"0{idx+1}", 0.95, y, 0.55, 0.28, size=14, bold=True, color=[TEAL, BLUE, PINK, ORANGE, GREEN][idx])
        add_text(slide, head, 1.65, y, 2.4, 0.28, size=15, bold=True, color=NAVY)
        add_text(slide, body, 4.25, y + 0.02, 7.1, 0.26, size=13, color=GRAY)
    add_text(slide, "最适合强调的差异化：我们把 RAG 的关键点前移到了“高质量文档入库”。", 1.0, 6.25, 11.0, 0.34, size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    # 15 roadmap
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "后续规划：从轻量调度升级为 LangGraph Agent", "当前系统已经能演示主链路，后续重点是节点化、可恢复、可扩展。", "15 规划")
    panel(slide, 0.8, 2.0, 5.55, 3.8, fill=WHITE)
    add_text(slide, "Ingestion Agent", 1.15, 2.28, 2.4, 0.3, size=17, bold=True, color=TEAL)
    add_multiline(slide, ["FileProbe", "StrategyPlanner", "Parse / OCR", "Normalize", "ChunkQuality", "Index", "QualityReport"], 1.18, 2.84, 4.0, 2.0, size=14, color=NAVY, bullet=True)
    panel(slide, 6.9, 2.0, 5.55, 3.8, fill=WHITE)
    add_text(slide, "QA Agent", 7.25, 2.28, 2.4, 0.3, size=17, bold=True, color=BLUE)
    add_multiline(slide, ["Permission", "QuestionClassifier", "RetrievalPlanner", "HybridRetrieve", "Rerank", "EvidenceCheck", "Answer / Log"], 7.28, 2.84, 4.0, 2.0, size=14, color=NAVY, bullet=True)
    add_text(slide, "目标：把文档入库和问答过程拆成可观察、可回滚、可扩展的 Agent 工作流。", 1.0, 6.32, 11.0, 0.34, size=18, bold=True, color=PINK, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_deck()
