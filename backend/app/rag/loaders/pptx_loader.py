import json
from pathlib import Path
from typing import Any, Iterable

from app.rag.loaders.base import BaseLoader
from app.rag.schemas import DocumentBlock, SourceLocator


class PptxLoader(BaseLoader):
    """Extract slide text, tables, and speaker notes from a PPTX file."""

    def load(self, path: str | Path) -> list[DocumentBlock]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError(
                "PPTX parsing requires python-pptx. "
                "Install: pip install -r requirements-rag.txt"
            ) from exc

        file_path = Path(path)
        presentation = Presentation(str(file_path))
        blocks: list[DocumentBlock] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            blocks.extend(
                _slide_blocks(
                    slide,
                    source=file_path.name,
                    slide_index=slide_index,
                )
            )
        return blocks


def _slide_blocks(
    slide: Any,
    *,
    source: str,
    slide_index: int,
) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    table_blocks: list[DocumentBlock] = []
    slide_title = _slide_title(slide)
    title_shape = getattr(slide.shapes, "title", None)
    title_shape_id = getattr(title_shape, "shape_id", None)
    text_items: list[tuple[int, int, int, str]] = []
    table_index = 0
    image_count = 0

    for shape_index, shape in enumerate(_iter_shapes(slide.shapes), start=1):
        if getattr(shape, "has_table", False):
            table_index += 1
            table_blocks.extend(
                _table_blocks(
                    shape.table,
                    source=source,
                    slide_index=slide_index,
                    slide_title=slide_title,
                    table_index=table_index,
                    shape_index=shape_index,
                )
            )
            continue
        if _is_picture(shape):
            image_count += 1
            continue
        if _is_ignored_placeholder(shape):
            continue
        if (
            title_shape_id is not None
            and getattr(shape, "shape_id", None) == title_shape_id
        ):
            continue
        text = _shape_text(shape)
        if text:
            text_items.append(
                (
                    _safe_int(getattr(shape, "top", 0)),
                    _safe_int(getattr(shape, "left", 0)),
                    shape_index,
                    text,
                )
            )

    text_items.sort(key=lambda item: (item[0], item[1], item[2]))
    text_parts = [text for _, _, _, text in text_items]
    if slide_title or text_parts:
        content = "\n".join(
            [
                *(["幻灯片标题：" + slide_title] if slide_title else []),
                *text_parts,
            ]
        )
        metadata: dict[str, Any] = {
            "slide_title": slide_title,
            "shape_count": len(text_items),
            "image_count": image_count,
            "table_count": table_index,
        }
        if text_items:
            metadata["shape_indices"] = ",".join(
                str(item[2]) for item in text_items
            )
        blocks.append(
            DocumentBlock(
                text=content,
                block_type="slide_text",
                source=source,
                document_format="pptx",
                extraction_method="python_pptx",
                locator=SourceLocator(
                    slide=slide_index,
                    section_path=slide_title or None,
                    block_index=text_items[0][2] if text_items else 1,
                ),
                metadata=metadata,
            )
        )

    blocks.extend(table_blocks)
    notes_text = _speaker_notes_text(slide)
    if notes_text:
        prefix = f"幻灯片标题：{slide_title}\n" if slide_title else ""
        blocks.append(
            DocumentBlock(
                text=prefix + notes_text,
                block_type="speaker_notes",
                source=source,
                document_format="pptx",
                extraction_method="python_pptx_notes",
                locator=SourceLocator(
                    slide=slide_index,
                    section_path=slide_title or None,
                ),
                metadata={"slide_title": slide_title},
            )
        )
    return blocks


def _iter_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        MSO_SHAPE_TYPE = None

    for shape in shapes:
        if (
            MSO_SHAPE_TYPE is not None
            and getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP
        ):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _slide_title(slide: Any) -> str:
    title_shape = getattr(slide.shapes, "title", None)
    return _shape_text(title_shape) if title_shape is not None else ""


def _shape_text(shape: Any) -> str:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return ""
    lines: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        text = _clean_text(paragraph.text)
        if not text:
            continue
        level = max(0, _safe_int(getattr(paragraph, "level", 0)))
        prefix = ("  " * level + "- ") if level > 0 else ""
        lines.append(prefix + text)
    return "\n".join(lines).strip()


def _speaker_notes_text(slide: Any) -> str:
    if not bool(getattr(slide, "has_notes_slide", False)):
        return ""
    notes_slide = slide.notes_slide
    text_frame = getattr(notes_slide, "notes_text_frame", None)
    if text_frame is None:
        return ""
    return "\n".join(
        text
        for text in (_clean_text(paragraph.text) for paragraph in text_frame.paragraphs)
        if text
    ).strip()


def _table_blocks(
    table: Any,
    *,
    source: str,
    slide_index: int,
    slide_title: str,
    table_index: int,
    shape_index: int,
) -> list[DocumentBlock]:
    rows = [
        [_clean_text(cell.text) for cell in row.cells]
        for row in table.rows
    ]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []

    column_count = max(len(row) for row in rows)
    rows = [_pad_row(row, column_count) for row in rows]
    headers = _dedupe_headers(rows[0])
    data_rows = rows[1:]
    table_id = f"pptx:s{slide_index}:t{table_index}"
    common_metadata = {
        "slide_title": slide_title,
        "table_id": table_id,
        "table_title": slide_title or None,
        "table_row_count": len(data_rows),
        "table_column_count": len(headers),
        "table_columns": " | ".join(headers),
        "table_chunking_strategy": "row_aware",
        "shape_index": shape_index,
    }
    blocks = [
        DocumentBlock(
            text=_table_markdown(headers, data_rows),
            block_type="table_markdown",
            source=source,
            document_format="pptx",
            extraction_method="python_pptx_table",
            locator=SourceLocator(
                slide=slide_index,
                section_path=slide_title or None,
                table_index=table_index,
                block_index=shape_index,
            ),
            metadata=common_metadata,
        )
    ]

    for row_index, row in enumerate(data_rows, start=1):
        row_map = dict(zip(headers, row, strict=False))
        facts = "；".join(
            f"{key}：{value}"
            for key, value in row_map.items()
            if value
        )
        if not facts:
            continue
        title_prefix = f"{slide_title}。" if slide_title else ""
        blocks.append(
            DocumentBlock(
                text=f"{title_prefix}表格第{row_index}行：{facts}。",
                block_type="table_row_fact",
                source=source,
                document_format="pptx",
                extraction_method="python_pptx_table",
                locator=SourceLocator(
                    slide=slide_index,
                    section_path=slide_title or None,
                    table_index=table_index,
                    block_index=row_index,
                ),
                metadata={
                    **common_metadata,
                    "table_row_index": row_index,
                    "table_row_fields_json": json.dumps(
                        row_map,
                        ensure_ascii=False,
                    ),
                    "table_chunking_strategy": "row_fact",
                },
            )
        )
    return blocks


def _is_picture(shape: Any) -> bool:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return False
    return getattr(shape, "shape_type", None) in {
        MSO_SHAPE_TYPE.PICTURE,
        MSO_SHAPE_TYPE.LINKED_PICTURE,
    }


def _is_ignored_placeholder(shape: Any) -> bool:
    if not bool(getattr(shape, "is_placeholder", False)):
        return False
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER
    except ImportError:
        return False
    placeholder_type = getattr(shape.placeholder_format, "type", None)
    return placeholder_type in {
        PP_PLACEHOLDER.DATE,
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.SLIDE_NUMBER,
    }


def _table_markdown(headers: list[str], rows: list[list[str]]) -> str:
    header_line = "| " + " | ".join(headers) + " |"
    separator = "| " + " | ".join("---" for _ in headers) + " |"
    body = ["| " + " | ".join(row) + " |" for row in rows]
    return "\n".join([header_line, separator, *body])


def _dedupe_headers(values: list[str]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = value or f"列{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _pad_row(row: list[str], column_count: int) -> list[str]:
    return [*row, *([""] * (column_count - len(row)))]


def _clean_text(value: Any) -> str:
    lines = [" ".join(line.split()) for line in str(value or "").splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _safe_int(value: Any) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0
