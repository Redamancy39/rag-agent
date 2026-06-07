import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from app.rag.loaders.base import BaseLoader
from app.rag.schemas import DocumentBlock, SourceLocator

MAX_SCAN_ROWS = 20000
MAX_SCAN_COLUMNS = 256
MARKDOWN_ROWS_PER_BLOCK = 40
HEADER_SCAN_ROWS = 10


class XlsxLoader(BaseLoader):
    """Extract visible worksheet regions into structured DocumentBlock objects."""

    def load(self, path: str | Path) -> list[DocumentBlock]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "XLSX parsing requires openpyxl. "
                "Install: pip install -r requirements-rag.txt"
            ) from exc

        file_path = Path(path)
        formula_book = load_workbook(
            str(file_path),
            data_only=False,
            read_only=False,
        )
        value_book = load_workbook(
            str(file_path),
            data_only=True,
            read_only=False,
        )
        blocks: list[DocumentBlock] = []
        try:
            for formula_sheet in formula_book.worksheets:
                if formula_sheet.sheet_state != "visible":
                    continue
                value_sheet = value_book[formula_sheet.title]
                blocks.extend(
                    _worksheet_blocks(
                        formula_sheet,
                        value_sheet,
                        source=file_path.name,
                    )
                )
        finally:
            formula_book.close()
            value_book.close()
        return blocks


def _worksheet_blocks(
    formula_sheet: Any,
    value_sheet: Any,
    *,
    source: str,
) -> list[DocumentBlock]:
    from openpyxl.utils import get_column_letter

    max_row = min(int(formula_sheet.max_row or 0), MAX_SCAN_ROWS)
    max_column = min(int(formula_sheet.max_column or 0), MAX_SCAN_COLUMNS)
    if max_row <= 0 or max_column <= 0:
        return []

    merged_values = _merged_cell_values(formula_sheet, value_sheet)
    values: dict[tuple[int, int], str] = {}
    formulas: set[tuple[int, int]] = set()
    nonempty_by_row: dict[int, set[int]] = {}
    for row_index in range(1, max_row + 1):
        for column_index in range(1, max_column + 1):
            formula_cell = formula_sheet.cell(row_index, column_index)
            value_cell = value_sheet.cell(row_index, column_index)
            value, is_formula = _cell_text(
                formula_cell,
                value_cell,
                merged_values,
            )
            if not value:
                continue
            values[(row_index, column_index)] = value
            nonempty_by_row.setdefault(row_index, set()).add(column_index)
            if is_formula:
                formulas.add((row_index, column_index))

    if not values:
        return []

    blocks: list[DocumentBlock] = []
    table_index = 0
    for row_start, row_end in _contiguous_groups(sorted(nonempty_by_row)):
        active_columns = sorted(
            {
                column
                for row in range(row_start, row_end + 1)
                for column in nonempty_by_row.get(row, set())
            }
        )
        for column_start, column_end in _contiguous_groups(active_columns):
            if not _region_has_values(
                values,
                row_start,
                row_end,
                column_start,
                column_end,
            ):
                continue
            table_index += 1
            region_blocks = _region_blocks(
                formula_sheet,
                values=values,
                formulas=formulas,
                source=source,
                sheet_name=formula_sheet.title,
                sheet_state=formula_sheet.sheet_state,
                row_start=row_start,
                row_end=row_end,
                column_start=column_start,
                column_end=column_end,
                table_index=table_index,
            )
            blocks.extend(region_blocks)

    if max_row >= MAX_SCAN_ROWS or max_column >= MAX_SCAN_COLUMNS:
        blocks.append(
            DocumentBlock(
                text=(
                    "工作表数据范围超过当前安全扫描上限，"
                    f"仅处理前 {MAX_SCAN_ROWS} 行、前 {MAX_SCAN_COLUMNS} 列。"
                ),
                block_type="sheet_processing_notice",
                source=source,
                document_format="xlsx",
                extraction_method="openpyxl",
                locator=SourceLocator(
                    sheet=formula_sheet.title,
                    cell_range=(
                        f"A1:{get_column_letter(max_column)}{max_row}"
                    ),
                ),
                metadata={"scan_truncated": True},
            )
        )
    return blocks


def _region_blocks(
    worksheet: Any,
    *,
    values: dict[tuple[int, int], str],
    formulas: set[tuple[int, int]],
    source: str,
    sheet_name: str,
    sheet_state: str,
    row_start: int,
    row_end: int,
    column_start: int,
    column_end: int,
    table_index: int,
) -> list[DocumentBlock]:
    from openpyxl.utils import get_column_letter

    matrix = [
        [
            values.get((row_index, column_index), "")
            for column_index in range(column_start, column_end + 1)
        ]
        for row_index in range(row_start, row_end + 1)
    ]
    header_offset = _find_header_offset(
        worksheet,
        matrix,
        row_start=row_start,
        column_start=column_start,
    )
    header_row = row_start + header_offset
    preamble_rows = matrix[:header_offset]
    header_values = matrix[header_offset]
    data_rows = [
        (row_start + offset, row)
        for offset, row in enumerate(matrix[header_offset + 1 :], start=header_offset + 1)
        if any(row)
    ]
    region_range = (
        f"{get_column_letter(column_start)}{row_start}:"
        f"{get_column_letter(column_end)}{row_end}"
    )
    table_title = _table_title(preamble_rows) or sheet_name
    common_metadata = {
        "sheet_state": sheet_state,
        "table_id": f"xlsx:{sheet_name}:t{table_index}",
        "table_title": table_title,
        "table_header_row": header_row,
        "table_row_count": len(data_rows),
        "table_column_count": len(header_values),
        "table_region": region_range,
        "table_chunking_strategy": "row_aware",
    }
    blocks: list[DocumentBlock] = []

    preamble_text = _preamble_text(preamble_rows)
    if preamble_text:
        preamble_end = header_row - 1
        blocks.append(
            DocumentBlock(
                text=preamble_text,
                block_type="sheet_context",
                source=source,
                document_format="xlsx",
                extraction_method="openpyxl",
                locator=SourceLocator(
                    sheet=sheet_name,
                    cell_range=(
                        f"{get_column_letter(column_start)}{row_start}:"
                        f"{get_column_letter(column_end)}{preamble_end}"
                    ),
                    table_index=table_index,
                ),
                metadata=common_metadata,
            )
        )

    if not data_rows:
        text = "；".join(value for value in header_values if value)
        if text and not preamble_text:
            blocks.append(
                DocumentBlock(
                    text=text,
                    block_type="sheet_context",
                    source=source,
                    document_format="xlsx",
                    extraction_method="openpyxl",
                    locator=SourceLocator(
                        sheet=sheet_name,
                        cell_range=region_range,
                        table_index=table_index,
                    ),
                    metadata=common_metadata,
                )
            )
        return blocks

    headers = _dedupe_headers(header_values)
    common_metadata["table_columns"] = " | ".join(headers)
    for segment_index, segment_start in enumerate(
        range(0, len(data_rows), MARKDOWN_ROWS_PER_BLOCK),
        start=1,
    ):
        segment = data_rows[
            segment_start : segment_start + MARKDOWN_ROWS_PER_BLOCK
        ]
        segment_end_row = segment[-1][0]
        cell_range = (
            f"{get_column_letter(column_start)}{header_row}:"
            f"{get_column_letter(column_end)}{segment_end_row}"
        )
        blocks.append(
            DocumentBlock(
                text=_table_markdown(headers, [row for _, row in segment]),
                block_type="table_markdown",
                source=source,
                document_format="xlsx",
                extraction_method="openpyxl_table",
                locator=SourceLocator(
                    sheet=sheet_name,
                    cell_range=cell_range,
                    table_index=table_index,
                    block_index=segment_index,
                ),
                metadata={
                    **common_metadata,
                    "table_segment_index": segment_index,
                    "table_segment_count": (
                        len(data_rows) + MARKDOWN_ROWS_PER_BLOCK - 1
                    )
                    // MARKDOWN_ROWS_PER_BLOCK,
                },
            )
        )

    normalized_headers = [header.strip().lower() for header in headers]
    for row_index, row in data_rows:
        if [value.strip().lower() for value in row] == normalized_headers:
            continue
        row_map = dict(zip(headers, row, strict=False))
        facts = "；".join(
            f"{key}：{value}"
            for key, value in row_map.items()
            if value
        )
        if not facts:
            continue
        row_range = (
            f"{get_column_letter(column_start)}{row_index}:"
            f"{get_column_letter(column_end)}{row_index}"
        )
        formula_count = sum(
            1
            for column_index in range(column_start, column_end + 1)
            if (row_index, column_index) in formulas
        )
        blocks.append(
            DocumentBlock(
                text=f"{table_title}。第 {row_index} 行：{facts}。",
                block_type="table_row_fact",
                source=source,
                document_format="xlsx",
                extraction_method="openpyxl_table",
                locator=SourceLocator(
                    sheet=sheet_name,
                    cell_range=row_range,
                    table_index=table_index,
                    block_index=row_index,
                ),
                metadata={
                    **common_metadata,
                    "table_row_index": row_index,
                    "table_row_fields_json": json.dumps(
                        row_map,
                        ensure_ascii=False,
                    ),
                    "formula_count": formula_count,
                    "table_chunking_strategy": "row_fact",
                },
            )
        )
    return blocks


def _find_header_offset(
    worksheet: Any,
    matrix: list[list[str]],
    *,
    row_start: int,
    column_start: int,
) -> int:
    if len(matrix) <= 1:
        return 0

    width = max(1, len(matrix[0]))
    scores: list[tuple[float, int]] = []
    for offset, row in enumerate(matrix[:HEADER_SCAN_ROWS]):
        nonempty = [value for value in row if value]
        if not nonempty:
            continue
        density = len(nonempty) / width
        text_ratio = sum(not _looks_numeric(value) for value in nonempty) / len(nonempty)
        unique_ratio = len(set(nonempty)) / len(nonempty)
        styled = sum(
            _is_header_styled(
                worksheet.cell(row_start + offset, column_start + index)
            )
            for index, value in enumerate(row)
            if value
        ) / len(nonempty)
        next_density = 0.0
        if offset + 1 < len(matrix):
            next_density = sum(bool(value) for value in matrix[offset + 1]) / width
        score = (
            density * 2.0
            + text_ratio
            + unique_ratio * 0.35
            + styled * 1.4
            + next_density * 0.35
            - (1.0 - unique_ratio) * 2.0
        )
        if len(nonempty) == 1 and width > 1:
            score -= 1.4
        scores.append((score, offset))
    if not scores:
        return 0
    best_score = max(score for score, _ in scores)
    return min(offset for score, offset in scores if score >= best_score - 0.15)


def _merged_cell_values(
    formula_sheet: Any,
    value_sheet: Any,
) -> dict[tuple[int, int], tuple[str, bool]]:
    values: dict[tuple[int, int], tuple[str, bool]] = {}
    for merged_range in formula_sheet.merged_cells.ranges:
        min_column, min_row, max_column, max_row = merged_range.bounds
        formula_cell = formula_sheet.cell(min_row, min_column)
        value_cell = value_sheet.cell(min_row, min_column)
        text, is_formula = _raw_cell_text(formula_cell, value_cell)
        if not text:
            continue
        if (max_row - min_row + 1) * (max_column - min_column + 1) > 10000:
            continue
        for row_index in range(min_row, max_row + 1):
            for column_index in range(min_column, max_column + 1):
                values[(row_index, column_index)] = (text, is_formula)
    return values


def _cell_text(
    formula_cell: Any,
    value_cell: Any,
    merged_values: dict[tuple[int, int], tuple[str, bool]],
) -> tuple[str, bool]:
    merged = merged_values.get((formula_cell.row, formula_cell.column))
    if merged is not None:
        return merged
    return _raw_cell_text(formula_cell, value_cell)


def _raw_cell_text(formula_cell: Any, value_cell: Any) -> tuple[str, bool]:
    is_formula = formula_cell.data_type == "f"
    value = value_cell.value if is_formula and value_cell.value is not None else formula_cell.value
    return _display_value(value, formula_cell.number_format), is_formula


def _display_value(value: Any, number_format: str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        if number_format and "%" in number_format:
            return f"{value * 100:g}%"
        return f"{value:g}"
    return " ".join(str(value).replace("\r", "\n").split())


def _contiguous_groups(values: list[int]) -> list[tuple[int, int]]:
    if not values:
        return []
    groups: list[tuple[int, int]] = []
    start = previous = values[0]
    for value in values[1:]:
        if value != previous + 1:
            groups.append((start, previous))
            start = value
        previous = value
    groups.append((start, previous))
    return groups


def _region_has_values(
    values: dict[tuple[int, int], str],
    row_start: int,
    row_end: int,
    column_start: int,
    column_end: int,
) -> bool:
    return any(
        (row_index, column_index) in values
        for row_index in range(row_start, row_end + 1)
        for column_index in range(column_start, column_end + 1)
    )


def _table_title(rows: list[list[str]]) -> str:
    for row in reversed(rows):
        text = " ".join(_unique_row_values(row)).strip()
        if text:
            return text[:200]
    return ""


def _preamble_text(rows: list[list[str]]) -> str:
    lines = [
        "；".join(_unique_row_values(row))
        for row in rows
        if any(row)
    ]
    return "\n".join(line for line in lines if line).strip()


def _unique_row_values(row: list[str]) -> list[str]:
    values: list[str] = []
    previous = ""
    for value in row:
        if value and value != previous:
            values.append(value)
        previous = value
    return values


def _table_markdown(headers: list[str], rows: list[list[str]]) -> str:
    escaped_headers = [_markdown_value(value) for value in headers]
    header_line = "| " + " | ".join(escaped_headers) + " |"
    separator = "| " + " | ".join("---" for _ in headers) + " |"
    body = [
        "| " + " | ".join(_markdown_value(value) for value in row) + " |"
        for row in rows
    ]
    return "\n".join([header_line, separator, *body])


def _markdown_value(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def _dedupe_headers(values: list[str]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = value or f"列{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _is_header_styled(cell: Any) -> bool:
    font = getattr(cell, "font", None)
    fill = getattr(cell, "fill", None)
    return bool(
        getattr(font, "bold", False)
        or (
            fill is not None
            and getattr(fill, "fill_type", None) not in {None, "none"}
        )
    )


def _looks_numeric(value: str) -> bool:
    cleaned = value.strip().replace(",", "").replace("%", "")
    try:
        float(cleaned)
        return True
    except ValueError:
        return False
