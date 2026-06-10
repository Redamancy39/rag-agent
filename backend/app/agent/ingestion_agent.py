from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from app.core.config import get_settings
from app.core.logger import get_logger
from app.rag.llm_client import OpenAICompatibleLLM
from app.rag.text_quality import is_text_usable, text_quality_summary

logger = get_logger(__name__)

INGESTION_AGENT_PROMPT = """你是文档入库工具选择 Agent。你只能根据文件探测信息选择处理工具。

可选工具：
1. pymupdf_text_table：适合有正常 PDF 文本层、需要表格探测的 PDF。
2. pypdf_text：适合没有 PyMuPDF 时的普通文本层 PDF。
3. paddleocr：适合扫描件、文本层为空、乱码、不可读文本。
4. liteparse：适合旧版 DOC/PPT/XLS、图片密集 Office、原生解析失败或需要版面坐标的文档。
5. llm_table_normalizer：适合 OCR、LiteParse、表格多或需要把表格行整理成事实句的结果。

要求：
- 如果文本长度很低、readable=false、乱码比例高，应选择 OCR。
- DOC/PPT/XLS 必须选择 liteparse。
- 如果使用 OCR、LiteParse 或表格数量大于 0，建议启用 llm_table_normalizer。
- 不要建议不存在的工具。
- 只输出 JSON。

JSON:
{
  "primary_reader": "pymupdf_text_table",
  "use_ocr": false,
  "use_liteparse": false,
  "use_llm_preprocessing": true,
  "reason": "简短原因"
}
"""


@dataclass(slots=True)
class IngestionProbe:
    file_type: str
    file_size: int
    page_count: int | None = None
    sampled_pages: int = 0
    text_chars: int = 0
    table_count: int = 0
    readable: bool = False
    quality: dict[str, float | int] = field(default_factory=dict)
    sample_text: str = ""
    probe_error: str | None = None


@dataclass(slots=True)
class IngestionPlan:
    primary_reader: str
    use_ocr: bool
    use_liteparse: bool
    use_llm_preprocessing: bool
    reason: str
    probe: IngestionProbe
    source: str = "rules"

    @property
    def status_message(self) -> str:
        tools = [self.primary_reader]
        if self.use_ocr:
            tools.append("paddleocr")
        if self.use_liteparse and "liteparse" not in tools:
            tools.append("liteparse")
        if self.use_llm_preprocessing:
            tools.append("llm_table_normalizer")
        return f"入库 Agent 已选择工具：{', '.join(tools)}。原因：{self.reason}"


def plan_ingestion(path: str | Path, file_type: str) -> IngestionPlan:
    file_path = Path(path)
    probe = probe_document(file_path, file_type)
    return plan_from_probe(probe)


def plan_from_probe(probe: IngestionProbe) -> IngestionPlan:
    settings = get_settings()
    plan = _rule_based_plan(probe)
    needs_llm_planner = (
        probe.file_type == "pdf"
        or bool(probe.probe_error)
        or not probe.readable
    )
    if (
        settings.INGESTION_AGENT_LLM_ENABLED
        and settings.OPENAI_API_KEY
        and needs_llm_planner
    ):
        plan = _merge_llm_plan(probe, plan)
    return plan


def probe_to_dict(probe: IngestionProbe) -> dict[str, Any]:
    return asdict(probe)


def plan_to_dict(plan: IngestionPlan) -> dict[str, Any]:
    return {
        "primary_reader": plan.primary_reader,
        "use_ocr": plan.use_ocr,
        "use_liteparse": plan.use_liteparse,
        "use_llm_preprocessing": plan.use_llm_preprocessing,
        "reason": plan.reason,
        "source": plan.source,
    }


def probe_document(path: Path, file_type: str) -> IngestionProbe:
    settings = get_settings()
    probe = IngestionProbe(
        file_type=file_type,
        file_size=path.stat().st_size if path.exists() else 0,
    )
    if file_type == "docx":
        return _probe_docx(path, probe)
    if file_type == "pptx":
        return _probe_pptx(path, probe)
    if file_type == "xlsx":
        return _probe_xlsx(path, probe)
    if file_type != "pdf":
        probe.readable = True
        return probe

    try:
        import fitz
    except ImportError:
        probe.probe_error = "PyMuPDF is not installed"
        return probe

    try:
        sample_text_parts: list[str] = []
        with fitz.open(str(path)) as pdf:
            probe.page_count = pdf.page_count
            sample_pages = min(pdf.page_count, max(1, settings.INGESTION_AGENT_SAMPLE_PAGES))
            probe.sampled_pages = sample_pages
            for page_index in range(sample_pages):
                page = pdf[page_index]
                text = page.get_text("text") or ""
                sample_text_parts.append(text)
                probe.table_count += _safe_table_count(page)
        probe.sample_text = "\n".join(sample_text_parts).strip()[:1500]
        probe.text_chars = len(probe.sample_text)
        probe.quality = text_quality_summary(probe.sample_text)
        probe.readable = is_text_usable(probe.sample_text)
    except Exception as exc:
        probe.probe_error = str(exc)
    return probe


def _rule_based_plan(probe: IngestionProbe) -> IngestionPlan:
    settings = get_settings()
    if probe.file_type in {"doc", "ppt", "xls"}:
        return IngestionPlan(
            primary_reader="liteparse",
            use_ocr=False,
            use_liteparse=True,
            use_llm_preprocessing=settings.PREPROCESS_LLM_ENABLED,
            reason=(
                f"旧版 {probe.file_type.upper()} 使用 LiteParse 调用 LibreOffice "
                "转换并优先提取文本层；无可读文本时再自动尝试 OCR。"
                "此类转换结果优先进入大模型质量整理。"
            ),
            probe=probe,
        )
    if probe.file_type == "docx":
        if settings.LITEPARSE_ENABLED and (probe.probe_error or not probe.readable):
            return _liteparse_fallback_plan(probe, "DOCX 原生探测不可读或解析失败")
        use_llm = settings.PREPROCESS_LLM_ENABLED and probe.table_count > 0
        reason = "DOCX 使用 python-docx 提取标题、段落和表格。"
        if probe.table_count > 0:
            reason += f" 检测到 {probe.table_count} 个表格。"
        if use_llm:
            reason += " 已启用表格型切块的大模型整理。"
        return IngestionPlan(
            primary_reader="docx_loader",
            use_ocr=False,
            use_liteparse=False,
            use_llm_preprocessing=use_llm,
            reason=reason,
            probe=probe,
        )
    if probe.file_type == "pptx":
        if settings.LITEPARSE_ENABLED and (probe.probe_error or not probe.readable):
            return _liteparse_fallback_plan(probe, "PPTX 原生探测不可读或图片内容占主导")
        use_llm = settings.PREPROCESS_LLM_ENABLED and probe.table_count > 0
        reason = "PPTX 使用 python-pptx 提取幻灯片文本、表格和演讲者备注。"
        if probe.table_count > 0:
            reason += f" 检测到 {probe.table_count} 个表格。"
        if use_llm:
            reason += " 已启用表格型切块的大模型整理。"
        return IngestionPlan(
            primary_reader="pptx_loader",
            use_ocr=False,
            use_liteparse=False,
            use_llm_preprocessing=use_llm,
            reason=reason,
            probe=probe,
        )
    if probe.file_type == "xlsx":
        if settings.LITEPARSE_ENABLED and (probe.probe_error or not probe.readable):
            return _liteparse_fallback_plan(probe, "XLSX 原生探测不可读或图片内容占主导")
        return IngestionPlan(
            primary_reader="xlsx_loader",
            use_ocr=False,
            use_liteparse=False,
            use_llm_preprocessing=False,
            reason=(
                "XLSX 使用 openpyxl 确定性提取工作表、数据区域、表头和行事实；"
                "不执行逐行 LLM 预处理，避免大型表格产生过量调用。"
            ),
            probe=probe,
        )
    if probe.file_type != "pdf":
        return IngestionPlan(
            primary_reader=f"{probe.file_type}_loader",
            use_ocr=False,
            use_liteparse=False,
            use_llm_preprocessing=False,
            reason="非 PDF 文件，直接使用对应文本 loader。",
            probe=probe,
        )

    needs_ocr = (
        bool(probe.probe_error)
        or probe.text_chars < settings.OCR_MIN_EXTRACTED_CHARS
        or not probe.readable
    )
    use_ocr = settings.PREPROCESS_OCR_FALLBACK_ENABLED and needs_ocr
    use_llm = settings.PREPROCESS_LLM_ENABLED and (probe.table_count > 0 or use_ocr)
    if use_ocr:
        reason = "PDF 文本层为空、乱码或不可读，优先使用 OCR。"
        primary_reader = "paddleocr"
    elif probe.probe_error:
        reason = "PyMuPDF 探测失败，回退到 pypdf 文本读取。"
        primary_reader = "pypdf_text"
    else:
        reason = "PDF 文本层可读，优先使用 PyMuPDF 读取文本和表格。"
        primary_reader = "pymupdf_text_table"
    if probe.table_count > 0 and use_llm:
        reason += f" 探测到 {probe.table_count} 个表格，启用大模型表格事实整理。"
    return IngestionPlan(
        primary_reader=primary_reader,
        use_ocr=use_ocr,
        use_liteparse=False,
        use_llm_preprocessing=use_llm,
        reason=reason,
        probe=probe,
    )


def _merge_llm_plan(probe: IngestionProbe, fallback: IngestionPlan) -> IngestionPlan:
    try:
        payload = OpenAICompatibleLLM().complete_json(
            INGESTION_AGENT_PROMPT,
            _probe_prompt(probe, fallback),
        )
    except Exception as exc:
        logger.warning("Ingestion LLM planner failed, using rule plan: %s", exc)
        return fallback

    primary_reader = str(payload.get("primary_reader") or fallback.primary_reader)
    if primary_reader not in {
        "pymupdf_text_table",
        "pypdf_text",
        "paddleocr",
        "txt_loader",
        "markdown_loader",
        "docx_loader",
        "pptx_loader",
        "xlsx_loader",
        "liteparse",
    }:
        primary_reader = fallback.primary_reader

    llm_use_ocr = bool(payload.get("use_ocr", fallback.use_ocr))
    liteparse_capable = probe.file_type in {
        "pdf",
        "doc",
        "docx",
        "ppt",
        "pptx",
        "xls",
        "xlsx",
    }
    llm_use_liteparse = liteparse_capable and (
        bool(payload.get("use_liteparse", fallback.use_liteparse))
        or primary_reader == "liteparse"
        or (llm_use_ocr and probe.file_type != "pdf")
    )
    llm_use_llm = bool(payload.get("use_llm_preprocessing", fallback.use_llm_preprocessing))

    # Deterministic safety gates keep mandatory legacy conversion and OCR fallbacks enabled.
    use_ocr = fallback.use_ocr or (
        llm_use_ocr and (probe.file_type == "pdf" or llm_use_liteparse)
    )
    use_liteparse = fallback.use_liteparse or llm_use_liteparse
    use_llm = fallback.use_llm_preprocessing or llm_use_llm
    reason = str(payload.get("reason") or fallback.reason).strip()[:500]
    if not reason:
        reason = fallback.reason
    return IngestionPlan(
        primary_reader=(
            "liteparse"
            if use_liteparse
            else "paddleocr"
            if use_ocr and probe.file_type == "pdf"
            else primary_reader
        ),
        use_ocr=use_ocr,
        use_liteparse=use_liteparse,
        use_llm_preprocessing=use_llm,
        reason=reason,
        probe=probe,
        source="llm+rules",
    )


def _probe_prompt(probe: IngestionProbe, fallback: IngestionPlan) -> str:
    return (
        "文件探测信息：\n"
        f"- file_type: {probe.file_type}\n"
        f"- file_size: {probe.file_size}\n"
        f"- page_count: {probe.page_count}\n"
        f"- sampled_pages: {probe.sampled_pages}\n"
        f"- text_chars: {probe.text_chars}\n"
        f"- table_count: {probe.table_count}\n"
        f"- readable: {probe.readable}\n"
        f"- quality: {probe.quality}\n"
        f"- probe_error: {probe.probe_error}\n"
        f"- rule_plan: {fallback.primary_reader}, "
        f"use_ocr={fallback.use_ocr}, "
        f"use_liteparse={fallback.use_liteparse}, "
        f"use_llm_preprocessing={fallback.use_llm_preprocessing}, "
        f"reason={fallback.reason}\n"
        f"文本样本：\n{probe.sample_text[:800]}"
    )


def _liteparse_fallback_plan(
    probe: IngestionProbe,
    reason: str,
) -> IngestionPlan:
    settings = get_settings()
    return IngestionPlan(
        primary_reader="liteparse",
        use_ocr=settings.LITEPARSE_OCR_ENABLED,
        use_liteparse=True,
        use_llm_preprocessing=settings.PREPROCESS_LLM_ENABLED,
        reason=(
            f"{reason}，回退 LiteParse 进行格式转换、版面提取和选择性 OCR；"
            "解析结果优先进入大模型结构化整理。"
        ),
        probe=probe,
    )


def _safe_table_count(page: Any) -> int:
    find_tables = getattr(page, "find_tables", None)
    if find_tables is None:
        return 0
    try:
        table_finder = find_tables()
    except Exception:
        return 0
    return len(getattr(table_finder, "tables", []) or [])


def _probe_docx(path: Path, probe: IngestionProbe) -> IngestionProbe:
    try:
        from docx import Document
    except ImportError:
        probe.probe_error = "python-docx is not installed"
        return probe

    try:
        document = Document(str(path))
        paragraph_texts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        sample_parts = paragraph_texts[:40]
        for table in document.tables[:5]:
            for row in table.rows[:5]:
                sample_parts.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )
        probe.sample_text = "\n".join(sample_parts).strip()[:1500]
        probe.text_chars = len(probe.sample_text)
        probe.table_count = len(document.tables)
        probe.quality = text_quality_summary(probe.sample_text)
        probe.readable = is_text_usable(probe.sample_text)
        if not probe.sample_text:
            probe.probe_error = "DOCX contains no readable paragraphs or table text"
    except Exception as exc:
        probe.probe_error = str(exc)
    return probe


def _probe_pptx(path: Path, probe: IngestionProbe) -> IngestionProbe:
    try:
        from pptx import Presentation
    except ImportError:
        probe.probe_error = "python-pptx is not installed"
        return probe

    try:
        presentation = Presentation(str(path))
        probe.page_count = len(presentation.slides)
        sample_count = min(
            probe.page_count,
            max(1, get_settings().INGESTION_AGENT_SAMPLE_PAGES),
        )
        probe.sampled_pages = sample_count
        sample_parts: list[str] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            for shape in _iter_pptx_shapes(slide.shapes):
                if getattr(shape, "has_table", False):
                    probe.table_count += 1
                    if slide_index <= sample_count:
                        for row in list(shape.table.rows)[:5]:
                            sample_parts.append(
                                " | ".join(cell.text.strip() for cell in row.cells)
                            )
                elif (
                    slide_index <= sample_count
                    and getattr(shape, "has_text_frame", False)
                    and shape.text.strip()
                ):
                    sample_parts.append(shape.text.strip())
            if (
                slide_index <= sample_count
                and bool(getattr(slide, "has_notes_slide", False))
            ):
                notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
                if notes_frame is not None and notes_frame.text.strip():
                    sample_parts.append(notes_frame.text.strip())

        probe.sample_text = "\n".join(sample_parts).strip()[:1500]
        probe.text_chars = len(probe.sample_text)
        probe.quality = text_quality_summary(probe.sample_text)
        probe.readable = is_text_usable(probe.sample_text)
        if not probe.sample_text:
            probe.probe_error = "PPTX contains no readable text, table text, or notes"
    except Exception as exc:
        probe.probe_error = str(exc)
    return probe


def _iter_pptx_shapes(shapes: Any) -> Any:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        yield from shapes
        return

    for shape in shapes:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_pptx_shapes(shape.shapes)
        else:
            yield shape


def _probe_xlsx(path: Path, probe: IngestionProbe) -> IngestionProbe:
    try:
        from openpyxl import load_workbook
    except ImportError:
        probe.probe_error = "openpyxl is not installed"
        return probe

    workbook = None
    try:
        workbook = load_workbook(
            str(path),
            data_only=False,
            read_only=True,
        )
        visible_sheets = [
            sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"
        ]
        probe.page_count = len(visible_sheets)
        sample_count = min(
            len(visible_sheets),
            max(1, get_settings().INGESTION_AGENT_SAMPLE_PAGES),
        )
        probe.sampled_pages = sample_count
        sample_parts: list[str] = []
        table_count = 0
        for sheet in visible_sheets:
            nonempty_rows = 0
            for row_index, row in enumerate(
                sheet.iter_rows(
                    min_row=1,
                    max_row=min(int(sheet.max_row or 0), 40),
                    max_col=min(int(sheet.max_column or 0), 20),
                    values_only=True,
                ),
                start=1,
            ):
                values = [
                    " ".join(str(value).split())
                    for value in row
                    if value is not None and str(value).strip()
                ]
                if not values:
                    continue
                nonempty_rows += 1
                if sheet in visible_sheets[:sample_count]:
                    sample_parts.append(
                        f"{sheet.title}!{row_index}: " + " | ".join(values)
                    )
            if nonempty_rows >= 2:
                table_count += 1

        probe.table_count = table_count
        probe.sample_text = "\n".join(sample_parts).strip()[:1500]
        probe.text_chars = len(probe.sample_text)
        probe.quality = text_quality_summary(probe.sample_text)
        probe.readable = is_text_usable(probe.sample_text)
        if not probe.sample_text:
            probe.probe_error = "XLSX contains no readable cells in visible worksheets"
    except Exception as exc:
        probe.probe_error = str(exc)
    finally:
        if workbook is not None:
            workbook.close()
    return probe
